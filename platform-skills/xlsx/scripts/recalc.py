"""Recalculate every formula with LibreOffice and report error cells.

Empirically verified inside the sandbox image (`EXPERT_WORK_SKILLS_TEST_IMAGE`), as the
task brief requires before trusting this file:

(a) filter-name quoting for same-format `--convert-to`: `run_soffice()` builds argv as a
    Python list and execs it directly (no shell), so a filter argument must NOT carry
    embedded double quotes. `--convert-to 'xlsx:Calc MS Excel 2007 XML'` is accepted
    (soffice logs `using filter : Calc MS Excel 2007 XML`, exit 0, output written).
    `--convert-to 'xlsx:"Calc MS Excel 2007 XML"'` (quotes baked into the string, as a
    shell-quoting habit would produce) still exits 0 but soffice logs an
    `SfxBaseModel::impl_store ... failed: 0x81a (Io/Parameter)` error and writes nothing
    -- a silent failure this module must never emit. Also verified: contrary to this
    file's original assumption, soffice does not actually refuse a same-format
    `--convert-to` (`--convert-to xlsx` on an .xlsx input succeeds fine); the rejection
    lives in `_office.convert()`'s own early guard, which is why `_recalc_same_format()`
    calls `_office.run_soffice()` directly instead.
(b) `OOXMLRecalcMode`/`ODFRecalcMode=0` (the xcu `_office.py` writes when `recalc=True`):
    for a formula cell that has NO cached value at all (e.g. anything openpyxl wrote,
    since openpyxl never computes formulas), LibreOffice always computes one on load --
    with or without this xcu, `--convert-to` or not -- there is no stored value to keep,
    so it has no other choice. What the xcu actually controls is whether an EXISTING
    (possibly stale) cached value is trusted or recomputed: hand-patching a valid cached
    `<v>` next to an untouched `<f>` to a wrong number and round-tripping through
    `--convert-to` reproduces it directly -- recalc=False keeps the wrong hand-patched
    number, recalc=True overwrites it with the correct recomputed one (see
    `case_xlsx_recalc.py`'s "陈旧缓存" block, which drives this through this very script).
    recalc=True is kept unconditionally: harmless when there is nothing stale (per (b)),
    and the only lever this skill has against a stale cache in a real-world (non-openpyxl)
    input.
"""

from __future__ import annotations

import argparse
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import _office
import openpyxl
from _cli import emit_json, fail, resolve_io

_ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("input")
    ap.add_argument("output")
    ap.add_argument("--timeout", type=float, default=120.0)
    args = ap.parse_args()
    src, dst = resolve_io(args.input, args.output)
    tmp = Path(tempfile.mkdtemp(prefix="recalc-"))
    try:
        staged = tmp / "in" / f"book{src.suffix.lower()}"
        staged.parent.mkdir()
        shutil.copyfile(src, staged)
        try:
            if src.suffix.lower() == ".xlsx":
                out = _recalc_same_format(staged, tmp / "out", args.timeout)
            else:
                out = _office.convert(
                    staged, "xlsx", tmp / "out", timeout=args.timeout, recalc=True
                )
        except (_office.OfficeError, _office.OfficeTimeout) as exc:
            fail(str(exc))
        shutil.copyfile(out, dst)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    formulas_wb = openpyxl.load_workbook(dst)
    values_wb = openpyxl.load_workbook(dst, data_only=True)
    formulas, errors = 0, []
    for ws in formulas_wb.worksheets:
        vws = values_wb[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas += 1
                v = vws[cell.coordinate].value
                if isinstance(v, str) and v in _ERRORS:
                    errors.append({"sheet": ws.title, "cell": cell.coordinate, "value": v})
    emit_json({"formulas": formulas, "errors": errors})
    # 3, not 2: argparse already owns 2 (usage error), and 1 is "recalc itself failed" —
    # a caller keying on the exit code alone must be able to tell the three apart.
    sys.exit(3 if errors else 0)


def _recalc_same_format(src: Path, out_dir: Path, timeout: float) -> Path:
    """xlsx -> xlsx round-trip via the explicit Calc filter name (see module docstring (a))."""
    out_dir.mkdir(parents=True, exist_ok=True)
    proc = _office.run_soffice(
        ["--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out_dir), str(src)],
        timeout=timeout,
        recalc=True,
    )
    target = out_dir / f"{src.stem}.xlsx"
    if proc.returncode != 0 or not target.is_file():
        detail = proc.stderr.strip()[-500:]
        raise _office.OfficeError(f"重算失败（退出码 {proc.returncode}）：{detail}")  # noqa: RUF001
    return target


if __name__ == "__main__":
    main()
