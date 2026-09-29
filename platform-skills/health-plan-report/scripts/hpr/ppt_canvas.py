"""Drawing surface for the PPT writer: a slide canvas with theme-aware shapes and text, and the
page chrome (header band, footer) every content page carries."""

from __future__ import annotations

from pathlib import Path
from typing import Any, BinaryIO

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Pt

from hpr.icons import ICONS
from hpr.measure import LINE
from hpr.ppt_layout import (
    BODY_W,
    FOOTER_BOTTOM,
    MARGIN_X,
    SLIDE_W,
    Ctx,
    Page,
    brand_line,
    lh,
)
from hpr.theme import Theme, on_color

BODY, CHROME, DECO = "hpr:body", "hpr:chrome", "hpr:deco"

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}


def emu(pt: float) -> Emu:
    return Emu(round(pt * 12700))


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_[1:])


def _set_font(run: Any, t: Theme, size: float, color: str, bold: bool) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    run.font.name = t.font_latin
    rpr = run.font._element
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = OxmlElement("a:ea")
        rpr.insert_element_before(
            ea, "a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"
        )
    ea.set("typeface", t.font_cn)


def _drop_style(shp: Any) -> None:
    """Remove the theme style reference so no renderer adds the theme's shadow/effects."""
    style = shp._element.find(qn("p:style"))
    if style is not None:
        shp._element.remove(style)


class Canvas:
    def __init__(self, slide: Any, theme: Theme) -> None:
        self.slide = slide
        self.t = theme

    def rect(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        fill: str | None,
        *,
        line: str | None = None,
        rounded: bool = False,
        name: str = DECO,
    ) -> Any:
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
        shp = self.slide.shapes.add_shape(kind, emu(x), emu(y), emu(w), emu(h))
        shp.name = name
        if rounded:
            shp.adjustments[0] = min(0.5, 4.0 / max(1.0, min(w, h)))
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(0.75)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        _drop_style(shp)
        return shp

    def oval(
        self, x: float, y: float, w: float, h: float, fill: str | None, line: str | None = None
    ) -> Any:
        shp = self.slide.shapes.add_shape(MSO_SHAPE.OVAL, emu(x), emu(y), emu(w), emu(h))
        shp.name = DECO
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(0.75)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        _drop_style(shp)
        return shp

    def text(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        text: str,
        size: float,
        color: str,
        *,
        bold: bool = False,
        align: str = "left",
        name: str = BODY,
        link: str | None = None,
    ) -> Any:
        tb = self.slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
        tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for i, para_text in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = _ALIGN[align]
            p.line_spacing = Pt(size * LINE)
            r = p.add_run()
            r.text = para_text
            _set_font(r, self.t, size, color, bold)
            if link:
                r.hyperlink.address = link
        return tb

    def polyline(self, pts: list[tuple[float, float]], color: str, width: float) -> Any:
        (x0, y0), rest = pts[0], pts[1:]
        builder = self.slide.shapes.build_freeform(emu(x0), emu(y0), scale=1.0)
        builder.add_line_segments([(emu(x), emu(y)) for x, y in rest], close=False)
        shp = builder.convert_to_shape()
        shp.name = DECO
        shp.fill.background()
        shp.line.color.rgb = rgb(color)
        shp.line.width = Pt(width)
        shp.shadow.inherit = False
        _drop_style(shp)
        return shp

    def icon(self, name: str, x: float, y: float, size: float, color: str) -> None:
        k = size / 24.0
        for pl in ICONS[name]:
            self.polyline([(x + px * k, y + py * k) for px, py in pl], color, 1.5)

    def picture(
        self, src: Path | BinaryIO, x: float, y: float, w: float, h: float, name: str
    ) -> Any:
        image = str(src) if isinstance(src, Path) else src
        pic = self.slide.shapes.add_picture(image, emu(x), emu(y), emu(w), emu(h))
        pic.name = name
        return pic


# ---------- chrome ----------


def _footer(cv: Canvas, content: dict, style: dict, ctx: Ctx) -> None:
    t, m = ctx.theme, ctx.m
    brand = content.get("brand") or {}
    line1 = brand_line(brand)
    line1_h = m.lines(line1, BODY_W, t.caption) * lh(t.caption) if line1 else 0.0
    disc = brand.get("disclaimer", "")
    n_disc = m.lines(disc, BODY_W, t.caption) if disc else 0
    text_h = line1_h + n_disc * lh(t.caption)
    if not text_h:
        return
    top = FOOTER_BOTTOM - text_h
    cv.rect(MARGIN_X, top - t.gap_s, BODY_W, 0.75, t.line)
    align = style["footer.align"]
    y = top
    if line1:
        cv.text(MARGIN_X, y, BODY_W, line1_h, line1, t.caption, t.muted, align=align, name=CHROME)
        y += line1_h
    if disc:
        cv.text(
            MARGIN_X,
            y,
            BODY_W,
            n_disc * lh(t.caption),
            disc,
            t.caption,
            t.muted,
            align=align,
            name=CHROME,
        )


def _header(cv: Canvas, page: Page, number: str, style: dict, ctx: Ctx) -> None:
    t = ctx.theme
    cv.rect(0, 0, SLIDE_W * 0.6, 4, t.primary)
    cv.rect(SLIDE_W * 0.6, 0, SLIDE_W * 0.4, 4, t.accent)
    x = MARGIN_X
    if style["section.icons"]:
        cv.rect(MARGIN_X, 22, 28, 28, t.primary, rounded=True)
        cv.icon(page.icon, MARGIN_X + 5, 27, 18, on_color(t.primary))
        x = MARGIN_X + 40
    title_w = SLIDE_W - MARGIN_X - x - 100
    cv.text(
        x,
        36 - lh(t.title) / 2,
        title_w,
        lh(t.title),
        page.title,
        t.title,
        t.ink,
        bold=True,
        name=CHROME,
    )
    if style["footer.page_number"]:
        cv.text(
            SLIDE_W - MARGIN_X - 90,
            36 - lh(t.caption) / 2,
            90,
            lh(t.caption),
            number,
            t.caption,
            t.muted,
            align="right",
            name=DECO,
        )
    cv.rect(MARGIN_X, 62, BODY_W, 0.75, t.line)


def _bg(slide: Any, color: str) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)
