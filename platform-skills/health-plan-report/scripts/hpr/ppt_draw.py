"""Draw a paginated plan onto a 16:9 deck (python-pptx native objects, visual direction A)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Pt

from hpr.blocks import RenderError, section_prims
from hpr.common import clock_minutes
from hpr.images import IMAGE_ERRORS, image_blocks, strict_images
from hpr.measure import LINE, Measurer
from hpr.page_map import pptx_page_map
from hpr.ppt_canvas import BODY, CHROME, DECO, Canvas, _bg, _footer, _header, _set_font, emu, rgb
from hpr.ppt_charts import add_chart
from hpr.ppt_cover import _cover, _end, _show_toc, _toc, band_section, draw_band
from hpr.ppt_layout import (
    BAND_ID,
    DONUT_H,
    MARGIN_X,
    SLIDE_W,
    STRIPE,
    Ctx,
    LayoutError,
    Placed,
    callout_title_height,
    card_height,
    card_parts,
    column_height,
    column_title_height,
    columns_rows,
    donut_legend_heights,
    grid_col_w,
    grid_rows,
    image_size,
    kv_geometry,
    kv_rows,
    lh,
    make_ctx,
    media_geometry,
    paginate,
    step_height,
    table_geometry,
    timeline_rows,
)
from hpr.prims import (
    Bullets,
    Callout,
    CardGrid,
    Chart,
    Columns,
    Image,
    InfoBand,
    KeyValue,
    Media,
    Paragraph,
    SubHeading,
    Table,
    TimeBars,
    Timeline,
)
from hpr.theme import Theme, build_theme, donut_colors, on_color

NO_GRID_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
_TONE = {
    "within": "within_text",
    "out": "out_text",
    "alert": "alert_text",
    "info": "primary_text",
    "neutral": "muted",
}


def _check_titles(content: dict, ctx: Ctx, style: dict) -> None:
    t, m = ctx.theme, ctx.m
    title_w = SLIDE_W - 2 * MARGIN_X - 100 - (40 if style["section.icons"] else 0)
    for si, sec in enumerate(content["sections"]):
        if m.lines(sec["title"], title_w, t.title, True) > 1:
            raise LayoutError(f"sections[{si}].title", "章节标题过长，页标题只能一行，请缩短")  # noqa: RUF001


# ---------- primitives ----------


def _tone(t: Theme, tone: str) -> str:
    return getattr(t, _TONE.get(tone, "muted"))


def _draw_grid(cv: Canvas, pl: Placed, g: CardGrid, ctx: Ctx) -> None:
    t = ctx.theme
    cw = grid_col_w(g, pl.w, ctx)
    y = pl.y
    for row in grid_rows(g):
        rh = max(card_height(c, cw, ctx) for c in row)
        for i, card in enumerate(row):
            x = pl.x + i * (cw + t.gap_m)
            stripe = t.out if card.tag and card.tag.tone in ("out", "alert") else t.primary
            cv.rect(x, y, cw, rh, t.background, line=t.line)
            cv.rect(x, y, cw, STRIPE, stripe)
            cy = y + STRIPE + t.gap_m
            inner = cw - 2 * t.gap_m
            for part in card_parts(card, inner, ctx):
                if part.kind == "bar" and card.bar:
                    _range_bar(cv, x + t.gap_m, cy, inner, card.bar, t)
                else:
                    color = {"title": t.muted, "value": t.ink, "line": t.ink}.get(part.kind)
                    if part.kind == "tag" and card.tag:
                        color = _tone(t, card.tag.tone)
                    cv.text(
                        x + t.gap_m,
                        cy,
                        inner,
                        part.height,
                        part.text,
                        part.size,
                        color or t.ink,
                        bold=part.bold,
                    )
                cy += part.height + t.gap_xs
        y += rh + t.gap_m


def _range_bar(cv: Canvas, x: float, y: float, w: float, bar: Any, t: Theme) -> None:
    span = bar.high - bar.low or 1.0
    dmin, dmax = bar.low - 0.6 * span, bar.high + 0.6 * span
    pos = (min(max(bar.value, dmin), dmax) - dmin) / (dmax - dmin)
    cv.rect(x, y + 3, w, 4, t.pale)
    cv.rect(
        x + w * (bar.low - dmin) / (dmax - dmin), y + 3, w * span / (dmax - dmin), 4, t.primary_soft
    )
    cv.rect(x + w * pos - 1, y, 2, 10, t.ink)


def _cell_border_bottom(cell: Any, color: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcpr.find(qn(tag))
        if old is not None:
            tcpr.remove(old)
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = OxmlElement(tag)
        if tag == "a:lnB":
            ln.set("w", str(int(0.75 * 12700)))
            fill = OxmlElement("a:solidFill")
            clr = OxmlElement("a:srgbClr")
            clr.set("val", color[1:])
            fill.append(clr)
            ln.append(fill)
        else:
            ln.set("w", "0")
            ln.append(OxmlElement("a:noFill"))
        tcpr.insert(i, ln)


def _draw_table(cv: Canvas, pl: Placed, tb: Table, ctx: Ctx) -> None:
    t = ctx.theme
    geo = table_geometry(tb, pl.w, ctx)
    gf = cv.slide.shapes.add_table(
        len(tb.rows) + 1, len(tb.columns), emu(pl.x), emu(pl.y), emu(pl.w), emu(pl.h)
    )
    gf.name = BODY
    table = gf.table
    table.first_row = True
    table.horz_banding = False
    style_id = table._tbl.tblPr.find(qn("a:tableStyleId"))
    if style_id is None:
        style_id = OxmlElement("a:tableStyleId")
        table._tbl.tblPr.append(style_id)
    style_id.text = NO_GRID_STYLE
    for i, w in enumerate(geo.col_w):
        table.columns[i].width = emu(w)
    table.rows[0].height = emu(geo.header_h)
    for r, h in enumerate(geo.row_h, start=1):
        table.rows[r].height = emu(h)
    for r, cells in enumerate([tb.columns, *tb.rows]):
        for c, value in enumerate(cells):
            cell = table.cell(r, c)
            _cell_border_bottom(cell, t.line)
            cell.margin_left = cell.margin_right = emu(geo.pad_x)
            cell.margin_top = cell.margin_bottom = emu(geo.pad_y)
            cell.vertical_anchor = MSO_ANCHOR.TOP
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(t.pale)
            else:
                cell.fill.background()
            highlight = r > 0 and tb.highlight_col == c and value not in ("", "—")
            color = t.primary_text if r == 0 else (t.out_text if highlight else t.ink)
            for k, line in enumerate(value.split("\n")):
                para = cell.text_frame.paragraphs[0] if k == 0 else cell.text_frame.add_paragraph()
                para.line_spacing = Pt(t.small * LINE)
                run = para.add_run()
                run.text = line
                _set_font(run, t, t.small, color, r == 0 or highlight)


def _draw_kv(cv: Canvas, pl: Placed, kv: KeyValue, ctx: Ctx) -> None:
    t = ctx.theme
    pair_w, label_w, heights = kv_geometry(kv, pl.w, ctx)
    y = pl.y
    for row, rh in zip(kv_rows(kv), heights, strict=True):
        for i, (k, v) in enumerate(row):
            x = pl.x + i * (pair_w + t.gap_l)
            cv.text(x, y + 2, label_w, rh - t.gap_s, k, t.small, t.muted)
            cv.text(
                x + label_w + t.gap_s, y, pair_w - label_w - t.gap_s, rh - t.gap_s, v, t.body, t.ink
            )
            cv.rect(x, y + rh - t.gap_s / 2, pair_w, 0.5, t.line)
        y += rh


def _draw_callout(cv: Canvas, pl: Placed, co: Callout, ctx: Ctx) -> None:
    t, m = ctx.theme, ctx.m
    bg = {"alert": t.pale_alert, "warn": t.pale_out}.get(co.tone, t.pale)
    edge = {"alert": t.alert, "warn": t.out}.get(co.tone, t.primary)
    cv.rect(pl.x, pl.y, pl.w, pl.h, bg, rounded=True)
    cv.rect(pl.x, pl.y, 4, pl.h, edge)
    x, y, w = pl.x + 4 + t.gap_m, pl.y + t.gap_m, pl.w - 4 - 2 * t.gap_m
    if co.title:
        th = callout_title_height(co, w, ctx)
        title_color = {"alert": t.alert_text, "warn": t.out_text}.get(co.tone, t.primary_text)
        cv.text(x, y, w, th, co.title, t.body, title_color, bold=True)
        y += th + t.gap_xs
    cv.text(x, y, w, m.lines(co.text, w, t.body) * lh(t.body), co.text, t.body, t.ink)


def _draw_chart(cv: Canvas, pl: Placed, ch: Chart, ctx: Ctx) -> None:
    t = ctx.theme
    if ch.kind != "donut":
        add_chart(cv.slide, ch, pl.x, pl.y, pl.w, pl.h, t)
        return
    size = min(pl.h, DONUT_H)
    add_chart(cv.slide, ch, pl.x, pl.y, size, size, t)
    if ch.center:
        cv.text(
            pl.x,
            pl.y + size / 2 - lh(t.heading) / 2,
            size,
            lh(t.heading),
            ch.center,
            t.heading,
            t.ink,
            bold=True,
            align="center",
        )
    colors = donut_colors(t)
    heights = donut_legend_heights(ch, pl.w, ctx)
    lx = pl.x + size + t.gap_l
    ly = pl.y + max(0.0, (size - sum(heights) - t.gap_xs * len(heights)) / 2)
    for i, (line, h) in enumerate(zip(ch.legend, heights, strict=True)):
        cv.rect(lx, ly + lh(t.body) / 2 - 5, 10, 10, colors[i % len(colors)])
        cv.text(lx + 18, ly, pl.w - size - t.gap_l - 18, h, line, t.body, t.ink)
        ly += h + t.gap_xs


def _draw_timeline(cv: Canvas, pl: Placed, tl: Timeline, ctx: Ctx) -> None:
    t, m = ctx.theme, ctx.m
    y = pl.y
    for row in timeline_rows(tl):
        sw = (pl.w - (len(row) - 1) * t.gap_m) / len(row)
        rh = max(step_height(s, sw, ctx) for s in row)
        cv.rect(pl.x, y + 5.5, pl.w, 1, t.line)
        for i, step in enumerate(row):
            x = pl.x + i * (sw + t.gap_m)
            cv.oval(x, y, 12, 12, t.primary)
            ly = y + 12 + t.gap_s
            label_h = m.lines(step.label, sw, t.body, True) * lh(t.body)
            cv.text(x, ly, sw, label_h, step.label, t.body, t.ink, bold=True)
            if step.lines:
                by = ly + label_h + t.gap_xs
                cv.rect(x, by, sw, y + rh - by, t.pale, rounded=True)
                cy = by + t.gap_s
                for ln in step.lines:
                    h = m.lines(ln, sw - 2 * t.gap_s, t.small) * lh(t.small)
                    cv.text(x + t.gap_s, cy, sw - 2 * t.gap_s, h, ln, t.small, t.ink)
                    cy += h
        y += rh + t.gap_l


def _draw_columns(cv: Canvas, pl: Placed, cols: Columns, ctx: Ctx) -> None:
    t, m = ctx.theme, ctx.m
    y = pl.y
    for row in columns_rows(cols):
        cw = (pl.w - (len(row) - 1) * t.gap_m) / len(row)
        rh = max(column_height(c, cw, ctx) for c in row)
        for i, col in enumerate(row):
            x = pl.x + i * (cw + t.gap_m)
            color = _tone(t, col.tone) if col.tone != "neutral" else t.primary_text
            fill = getattr(t, col.tone) if col.tone in ("within", "out", "alert") else t.primary
            cv.rect(x, y, cw, rh, t.pale, rounded=True)
            cv.rect(x, y, cw, STRIPE, fill)
            cy = y + STRIPE + t.gap_m
            title_h = column_title_height(col, cw, ctx)
            cv.text(
                x + t.gap_m, cy, cw - 2 * t.gap_m, title_h, col.title, t.heading, color, bold=True
            )
            cy += title_h + t.gap_s
            inner = cw - 2 * t.gap_m
            for item in col.items:
                h = m.lines(item, inner - t.small, t.small) * lh(t.small)
                cv.text(
                    x + t.gap_m, cy, t.small, lh(t.small), "·", t.small, color, bold=True, name=DECO
                )
                cv.text(x + t.gap_m + t.small, cy, inner - t.small, h, item, t.small, t.ink)
                cy += h + t.gap_xs
        y += rh + t.gap_m


def _draw_media(
    cv: Canvas, pl: Placed, md: Media, ctx: Ctx, base_dir: Path, warnings: list[str]
) -> None:
    t, m = ctx.theme, ctx.m
    cv.rect(pl.x, pl.y, pl.w, pl.h, t.pale, rounded=True)
    geo = media_geometry(md, pl.w, ctx)
    tx, tw = pl.x + t.gap_m, geo.text_w
    cv.text(tx, pl.y + t.gap_m, tw, geo.name_h, md.name, t.heading, t.primary_text, bold=True)
    dy = pl.y + t.gap_m + geo.name_h + t.gap_xs
    cv.text(
        tx, dy, tw, m.lines(md.description, tw, t.body) * lh(t.body), md.description, t.body, t.ink
    )
    bx = pl.x + pl.w * 0.62
    bw, bh = geo.box_w, pl.h - 2 * t.gap_m
    by = pl.y + t.gap_m
    cv.rect(bx, by, bw, bh, t.background, line=t.line, rounded=True)
    label, link_h = geo.link_text, geo.link_h
    if md.media_path:
        path = Path(md.media_path)
        path = path if path.is_absolute() else base_dir / path
        movie_h = bh - link_h - 2 * t.gap_s
        if path.suffix.lower() == ".mp4" and path.is_file() and movie_h > 30:
            try:
                mv = cv.slide.shapes.add_movie(
                    str(path),
                    emu(bx + t.gap_s),
                    emu(by + t.gap_s),
                    emu(bw - 2 * t.gap_s),
                    emu(movie_h),
                    mime_type="video/mp4",
                )
                mv.name = "hpr:media"
            except Exception as exc:  # any embed failure degrades to the link
                warnings.append(f"视频未能嵌入，已改为可点击链接：{md.name}（{exc}）")  # noqa: RUF001
        else:
            warnings.append(f"视频文件不可用，已改为可点击链接：{md.name}")  # noqa: RUF001
    cv.text(
        bx + t.gap_s,
        by + bh - link_h - t.gap_s,
        bw - 2 * t.gap_s,
        link_h,
        label,
        t.body,
        t.primary_text,
        bold=True,
        name=CHROME,
        link=md.url,
    )


def _draw_image(cv: Canvas, pl: Placed, im: Image, ctx: Ctx) -> None:
    t, m = ctx.theme, ctx.m
    w, h = image_size(im, pl.w, ctx)
    path = Path(im.path)
    path = path if path.is_absolute() else ctx.base_dir / path
    try:
        with strict_images():
            cv.picture(path, pl.x, pl.y, w, h, "hpr:image")
    except IMAGE_ERRORS as exc:
        raise RenderError(pl.path, f"图片无法读取（文件损坏或格式不支持）：{exc}") from exc  # noqa: RUF001
    if im.caption:
        ch = m.lines(im.caption, pl.w, t.caption) * lh(t.caption)
        cv.text(pl.x, pl.y + h + t.gap_xs, pl.w, ch, im.caption, t.caption, t.muted)


def _draw_timebars(cv: Canvas, pl: Placed, tb: TimeBars, ctx: Ctx) -> None:
    t, m = ctx.theme, ctx.m
    start, span = 18 * 60, 18 * 60
    label_w = 56.0
    ax, aw = pl.x + label_w, pl.w - label_w
    y = pl.y
    for i, (label, bed, wake) in enumerate(tb.rows):
        rh = lh(t.body)
        cv.text(pl.x, y, label_w - 8, rh, label, t.body, t.muted, bold=True, name=DECO)
        b, w_ = clock_minutes(bed), clock_minutes(wake)
        if w_ <= b:
            w_ += 24 * 60
        x0 = ax + aw * max(0, min(span, b - start)) / span
        x1 = ax + aw * max(0, min(span, w_ - start)) / span
        color = t.primary if i == len(tb.rows) - 1 else t.primary_soft
        cv.rect(x0, y + 2, max(4.0, x1 - x0), rh - 4, color, rounded=True)
        text = f"{bed} – {wake}"  # noqa: RUF001
        tw = m.width(text, t.small, True) * 1.2 + 4
        inside = tw <= (x1 - x0) - 8
        tx = x0 + 6 if inside else min(x1 + 6, pl.x + pl.w - tw)
        cv.text(
            tx,
            y + (rh - lh(t.small)) / 2,
            tw,
            lh(t.small),
            text,
            t.small,
            on_color(color) if inside else t.ink,
            bold=True,
        )
        y += rh + t.gap_s
    for k in range(0, 7):
        mins = start + k * 180
        hh = (mins // 60) % 24
        tx = ax + aw * k / 6 - 16
        cv.text(
            max(pl.x, min(tx, pl.x + pl.w - 32)),
            y,
            32,
            lh(t.caption),
            f"{hh:02d}:00",
            t.caption,
            t.muted,
            align="center",
            name=DECO,
        )


def _draw(cv: Canvas, pl: Placed, ctx: Ctx, base_dir: Path, warnings: list[str]) -> None:
    t, m = ctx.theme, ctx.m
    p = pl.prim
    if isinstance(p, SubHeading):
        cv.rect(pl.x - 10, pl.y + lh(t.heading) * 0.2, 3, lh(t.heading) * 0.6, t.accent)
        cv.text(pl.x, pl.y, pl.w, pl.h - t.gap_xs, p.text, t.heading, t.ink, bold=True)
    elif isinstance(p, Paragraph):
        pad = t.gap_m if p.boxed else 0.0
        if p.boxed:
            cv.rect(pl.x, pl.y, pl.w, pl.h, t.pale, rounded=True)
        cv.text(pl.x + pad, pl.y + pad, pl.w - 2 * pad, pl.h - 2 * pad, p.text, t.body, t.ink)
    elif isinstance(p, Bullets):
        indent = t.body * 1.4
        y = pl.y
        for n, item in enumerate(p.items, start=1):
            h = m.lines(item, pl.w - indent, t.body) * lh(t.body)
            marker = {"numbers": f"{n}.", "checks": "✓"}.get(p.style, "•")
            cv.text(pl.x, y, indent, lh(t.body), marker, t.body, t.primary, bold=True, name=DECO)
            cv.text(pl.x + indent, y, pl.w - indent, h, item, t.body, t.ink)
            y += h + t.gap_xs
    elif isinstance(p, CardGrid):
        _draw_grid(cv, pl, p, ctx)
    elif isinstance(p, Table):
        _draw_table(cv, pl, p, ctx)
    elif isinstance(p, KeyValue):
        _draw_kv(cv, pl, p, ctx)
    elif isinstance(p, Callout):
        _draw_callout(cv, pl, p, ctx)
    elif isinstance(p, Chart):
        _draw_chart(cv, pl, p, ctx)
    elif isinstance(p, Timeline):
        _draw_timeline(cv, pl, p, ctx)
    elif isinstance(p, Columns):
        _draw_columns(cv, pl, p, ctx)
    elif isinstance(p, Media):
        _draw_media(cv, pl, p, ctx, base_dir, warnings)
    elif isinstance(p, Image):
        _draw_image(cv, pl, p, ctx)
    elif isinstance(p, TimeBars):
        _draw_timebars(cv, pl, p, ctx)
    elif isinstance(p, InfoBand):
        draw_band(cv, pl.x, pl.y, pl.w, p, ctx)
    else:
        raise TypeError(type(p).__name__)


def render_pptx(
    content: dict,
    style: dict,
    out_path: Path,
    base_dir: Path,
    m: Measurer,
    page_map_out: list[dict[str, Any]] | None = None,
) -> list[str]:
    """Write the deck; ``page_map_out``, when given, receives which slides each part is on."""
    theme = build_theme(style, "pptx")
    ctx = make_ctx(content, theme, m, base_dir)
    _check_titles(content, ctx, style)
    image_blocks(content, style, base_dir)  # every image decodes, or its block path is reported
    pages = paginate(band_section(content) + section_prims(content, style), ctx)
    warnings: list[str] = []
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    toc = _show_toc(content, style)
    end = bool((content.get("brand") or {}).get("org_name"))
    total = 1 + int(toc) + len(pages) + int(end)
    _cover(prs, content, style, ctx, base_dir, warnings)
    if toc:
        _toc(prs, content, style, ctx)
    first_number = 2 + int(toc)
    for i, page in enumerate(pages):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _bg(slide, theme.background)
        cv = Canvas(slide, theme)
        _header(cv, page, f"{first_number + i} / {total}", style, ctx)
        for pl in page.placed:
            _draw(cv, pl, ctx, base_dir, warnings)
        _footer(cv, content, style, ctx)
    if end:
        _end(prs, content, ctx)
    prs.save(str(out_path))
    if page_map_out is not None:
        page_map_out.extend(pptx_page_map(content, pages, BAND_ID, toc, end))
    return warnings
