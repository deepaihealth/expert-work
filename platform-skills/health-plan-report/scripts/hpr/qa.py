"""Hard QA gates: every required text present, nothing overflows its box, nothing outside the
page, readable contrast (spec §7.2)."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from hpr.measure import LINE, Measurer
from hpr.style import INK, contrast
from hpr.theme import Theme, on_color

_EMU_PER_PT = 12700
_TEXT_NAMES = ("hpr:body", "hpr:chrome")


def norm(s: str) -> str:
    return re.sub(r"\s+", "", s)


def missing_texts(required: list[str], corpus: str) -> list[str]:
    c = norm(corpus)
    out: list[str] = []
    for r in required:
        if norm(r) not in c and r not in out:
            out.append(r)
    return out


def _shape_texts(sh: Any) -> list[str]:
    if sh.has_text_frame:
        return [sh.text_frame.text]
    if getattr(sh, "has_table", False) and sh.has_table:
        return [cell.text for row in sh.table.rows for cell in row.cells]
    return []


def pptx_corpus(path: Path) -> str:
    from pptx import Presentation

    body: list[str] = []
    chrome: list[str] = []
    for slide in Presentation(str(path)).slides:
        for sh in slide.shapes:
            if sh.name == "hpr:body":
                body += _shape_texts(sh)
            elif sh.name == "hpr:chrome":
                chrome += _shape_texts(sh)
    return "".join(body) + "\x00" + "\x00".join(chrome)


def pptx_bounds(path: Path) -> list[str]:
    from pptx import Presentation

    prs = Presentation(str(path))
    sw, sh_ = prs.slide_width, prs.slide_height
    tol = _EMU_PER_PT
    issues = []
    for n, slide in enumerate(prs.slides, start=1):
        for sh in slide.shapes:
            if None in (sh.left, sh.top, sh.width, sh.height):
                continue
            if (
                sh.left < -tol
                or sh.top < -tol
                or sh.left + sh.width > sw + tol
                or sh.top + sh.height > sh_ + tol
            ):
                issues.append(f"第 {n} 页形状「{sh.name}」超出页面")
    return issues


def _para_need(para: Any, width_pt: float, m: Measurer) -> float:
    runs = list(para.runs)
    size = runs[0].font.size.pt if runs and runs[0].font.size else 14.0
    bold = bool(runs and runs[0].font.bold)
    return m.lines(para.text, width_pt, size, bold) * size * LINE


def pptx_overflow(path: Path, m: Measurer) -> list[str]:
    from pptx import Presentation

    issues = []
    for n, slide in enumerate(Presentation(str(path)).slides, start=1):
        for sh in slide.shapes:
            if sh.name not in _TEXT_NAMES:
                continue
            if sh.has_text_frame:
                need = sum(
                    _para_need(p, sh.width / _EMU_PER_PT, m) for p in sh.text_frame.paragraphs
                )
                if need > sh.height / _EMU_PER_PT + 0.5:
                    box = sh.height / _EMU_PER_PT
                    label = sh.text_frame.text[:16]
                    issues.append(f"第 {n} 页「{label}」需要 {need:.0f}pt，框高 {box:.0f}pt")  # noqa: RUF001
            elif getattr(sh, "has_table", False) and sh.has_table:
                tbl = sh.table
                for r, row in enumerate(tbl.rows):
                    for c, cell in enumerate(row.cells):
                        w = (
                            tbl.columns[c].width - cell.margin_left - cell.margin_right
                        ) / _EMU_PER_PT
                        need = sum(_para_need(p, w, m) for p in cell.text_frame.paragraphs)
                        need += (cell.margin_top + cell.margin_bottom) / _EMU_PER_PT
                        if need > row.height / _EMU_PER_PT + 0.5:
                            issues.append(f"第 {n} 页表格第 {r + 1} 行第 {c + 1} 列溢出")
    return issues


def contrast_report(t: Theme) -> dict[str, Any]:
    ratios = {
        "ink_on_background": round(contrast(INK, t.background), 2),
        "primary_on_background": round(contrast(t.primary, t.background), 2),
        "text_on_primary": round(contrast(on_color(t.primary), t.primary), 2),
    }
    ok = (
        ratios["ink_on_background"] >= 7
        and ratios["primary_on_background"] >= 4.5
        and ratios["text_on_primary"] >= 4.5
    )
    return {"ok": ok, **ratios}


def qa_pptx(path: Path, required: list[str], theme: Theme, m: Measurer) -> dict[str, Any]:
    from pptx import Presentation

    missing = missing_texts(required, pptx_corpus(path))
    overflow = pptx_overflow(path, m)
    bounds = pptx_bounds(path)
    cr = contrast_report(theme)
    passed = not missing and not overflow and not bounds and cr["ok"]
    return {
        "status": "passed" if passed else "failed",
        "slides": len(Presentation(str(path)).slides),
        "missing": missing,
        "overflow": overflow,
        "out_of_bounds": bounds,
        "contrast": cr,
    }
