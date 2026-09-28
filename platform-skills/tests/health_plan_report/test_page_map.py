"""Page map: which pages each part of the report lands on, so a visual check opens the right
page (live replays spent looks on pages the agent guessed wrong)."""

import re
from types import SimpleNamespace

import pytest
from hpr.measure import Measurer
from hpr.page_map import pdf_page_map
from hpr.pdf_html import build_html
from hpr.ppt_draw import render_pptx
from hpr.style import Layer, resolve
from hpr.theme import build_theme
from PIL import Image as PILImage
from pptx import Presentation

M = Measurer("/nonexistent")


def _slide_texts(path) -> list[str]:
    out = []
    for slide in Presentation(str(path)).slides:
        parts = []
        for sh in slide.shapes:
            if sh.has_text_frame:
                parts.append(sh.text_frame.text)
            elif getattr(sh, "has_table", False) and sh.has_table:
                parts += [c.text for r in sh.table.rows for c in r.cells]
        out.append(re.sub(r"\s+", "", "".join(parts)))
    return out


@pytest.mark.parametrize("scale", ["standard", "large"])
def test_pptx_map_points_at_the_slides_that_show_each_section(tmp_path, sample, scale):
    PILImage.new("RGB", (640, 360), "#DDEEEE").save(tmp_path / "trend-note.png")
    style = resolve([Layer("t", {"type.scale": scale})], sample).style
    got: list[dict] = []
    render_pptx(sample, style, tmp_path / "o.pptx", tmp_path, M, got)
    texts = _slide_texts(tmp_path / "o.pptx")
    by_title = {e["title"]: e["pages"] for e in got}
    assert by_title["封面"] == [1]
    assert "客户信息" in by_title
    assert [e["section"] for e in got if "section" in e] == [s["id"] for s in sample["sections"]]
    covered = set()
    for e in got:
        pages = e["pages"]
        assert pages == list(range(pages[0], pages[-1] + 1))  # contiguous, ascending
        assert 1 <= pages[0] and pages[-1] <= len(texts)
        covered.update(pages)
        if "section" in e:
            assert re.sub(r"\s+", "", e["title"]) in texts[pages[0] - 1]
    assert covered == set(range(1, len(texts) + 1))  # every slide is accounted for


def test_pptx_map_names_toc_and_end_slides(tmp_path, sample):
    PILImage.new("RGB", (640, 360), "#DDEEEE").save(tmp_path / "trend-note.png")
    style = resolve([Layer("t", {"toc": "on"})], sample).style
    got: list[dict] = []
    render_pptx(sample, style, tmp_path / "o.pptx", tmp_path, M, got)
    n = len(Presentation(str(tmp_path / "o.pptx")).slides)
    by_title = {e["title"]: e["pages"] for e in got}
    assert by_title["目录"] == [2]
    assert by_title["封底"] == [n]


def _doc(*pages: dict) -> SimpleNamespace:
    return SimpleNamespace(pages=[SimpleNamespace(anchors=a) for a in pages])


def test_pdf_map_starts_at_the_heading_and_ends_at_the_end_marker(sample):
    first, second = sample["sections"][0]["id"], sample["sections"][1]["id"]
    doc = _doc(
        {},
        {"head-client-info": 0, "end-client-info": 0, f"head-sec-{first}": 0},
        {},
        # `second` and its heading leave empty, anchored fragments at the foot of page 4 while
        # the heading shows on page 5 (seen in the real image): the map must say 5
        {f"end-sec-{first}": 0, f"sec-{second}": 0, f"head-sec-{second}": 0},
        {f"sec-{second}": 0, f"head-sec-{second}": 0, f"end-sec-{second}": 0},
    )
    got = pdf_page_map(sample, doc, "client-info")
    by_title = {e["title"]: e["pages"] for e in got}
    assert by_title["封面"] == [1]
    assert by_title["客户信息"] == [2]
    assert got[2] == {"section": first, "title": sample["sections"][0]["title"], "pages": [2, 3, 4]}
    assert got[3]["pages"] == [5]


def test_pdf_html_marks_every_section_end_and_the_toc(sample, tmp_path):
    style = resolve([Layer("t", {"toc": "on"})], sample).style
    html, _ = build_html(sample, style, build_theme(style, "pdf"), tmp_path)
    assert 'id="hpr-toc"' in html
    assert html.count('<div id="end-') == len(sample["sections"]) + 1  # + the client band
    for s in sample["sections"]:
        start = html.index(f'id="sec-{s["id"]}"')
        head = html.index(f'<h2 id="head-sec-{s["id"]}"')
        end = html.index(f'id="end-sec-{s["id"]}"')
        assert start < head < end < html.index("</section>", start)
