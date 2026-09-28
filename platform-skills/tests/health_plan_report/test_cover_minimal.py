"""Minimal cover + client-info band (live-fix defect 1 and the user's cover ruling).

The cover shows only: LOGO / org name, title, subtitle, client name and one meta line
(generated date · manager). client.facts, period and data_basis move, in order, to the
「客户信息」 band at the top of the first body page. Covers fit at every type scale."""

import json
import re

import pytest
from hpr.common import band_columns, client_band_items, cover_meta_line, nobreak_tokens
from hpr.content import required_texts
from hpr.measure import SAFETY, Measurer
from hpr.pdf_html import build_html
from hpr.ppt_draw import render_pptx
from hpr.ppt_layout import BODY_W, band_geometry, lh, make_ctx
from hpr.prims import InfoBand
from hpr.qa import qa_pptx
from hpr.style import Layer, resolve
from hpr.theme import build_theme
from PIL import Image as PILImage
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

M = Measurer("/nonexistent")
SCALES = ("compact", "standard", "large")
VARIANTS = ("band", "split", "minimal")
TITLE_30 = "王女士二〇二六年秋季控糖减重与睡眠改善综合健康管理方案"[:30]
SUB_40 = "第 1 阶段（前 2 周）控糖与减重执行方案，配合饮食记录与每周复盘调整计划安排"  # noqa: RUF001
LONG_FACTS = [
    {"label": "性别 / 年龄", "value": "女 / 52 岁（1974-03）"},  # noqa: RUF001
    {"label": "身高 / 体重", "value": "160 cm / 68.4 kg（9-26 晨称）"},  # noqa: RUF001
    {"label": "健康状态", "value": "糖前期（2023-04 发现，病程 3 年余）"},  # noqa: RUF001
    {"label": "管理方向", "value": "控糖 + 减重"},
    {"label": "饮食限制", "value": "乳糖不耐受（少量酸奶可以），晚餐清淡"},  # noqa: RUF001
    {"label": "用药情况", "value": "暂无"},
    {"label": "睡眠", "value": "入睡困难，平均 6 小时"},  # noqa: RUF001
    {"label": "运动基础", "value": "每周快走 2 次，每次约 20 分钟"},  # noqa: RUF001
    {"label": "家族史", "value": "母亲 2 型糖尿病"},
    {"label": "数据截至", "value": "2026-09-28"},
]


@pytest.fixture
def base(tmp_path):
    PILImage.new("RGB", (640, 360), "#DDEEEE").save(tmp_path / "trend-note.png")
    PILImage.new("RGB", (300, 120), "#0B4F5C").save(tmp_path / "logo.png")
    return tmp_path


def _crowded(sample):
    sample["title"] = TITLE_30
    sample["subtitle"] = SUB_40
    sample["client"]["facts"] = json.loads(json.dumps(LONG_FACTS))
    sample["brand"]["logo_path"] = "logo.png"
    return sample


def _pptx(sample, base, layers=()):
    style = resolve(list(layers), sample).style
    out = base / "c.pptx"
    render_pptx(sample, style, out, base, M)
    return out, style


def _texts(slide) -> list[str]:
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text]


def test_cover_meta_line_and_band_items(sample):
    assert cover_meta_line(sample) == "2026-09-28 · 健康管理师 王老师"
    del sample["manager"]["title"]
    assert cover_meta_line(sample) == "2026-09-28 · 健康管理师 王老师"
    del sample["manager"]
    assert cover_meta_line(sample) == "2026-09-28"
    items = client_band_items(sample)
    facts = [(f["label"], f["value"]) for f in sample["client"]["facts"]]
    assert items == [*facts, ("阶段", "第 1 阶段"), ("数据依据", "近 30 天记录")]


@pytest.mark.parametrize("variant", VARIANTS)
def test_ppt_cover_holds_exactly_the_cover_elements(sample, base, variant):
    sample["brand"]["logo_path"] = "logo.png"
    out, _ = _pptx(sample, base, [Layer("x", {"cover.variant": variant})])
    cover = Presentation(str(out)).slides[0]
    want = {
        sample["brand"]["org_name"],
        sample["title"],
        sample["subtitle"],
        "客户",
        sample["client"]["name"],
        cover_meta_line(sample),
    }
    assert set(_texts(cover)) == want
    assert [sh.name for sh in cover.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE] == [
        "hpr:logo"
    ]


@pytest.mark.parametrize("variant", VARIANTS)
def test_ppt_cover_has_no_trend_line_decoration(sample, base, variant):
    out, _ = _pptx(sample, base, [Layer("x", {"cover.variant": variant})])
    cover = Presentation(str(out)).slides[0]
    kinds = {sh.shape_type for sh in cover.shapes}
    assert MSO_SHAPE_TYPE.FREEFORM not in kinds and MSO_SHAPE_TYPE.LINE not in kinds


def _band_slide(prs):
    return next(s for s in list(prs.slides)[1:] if "客户信息" in _texts(s))


def test_ppt_band_holds_every_fact_in_order_on_the_first_body_page(sample, base):
    out, _ = _pptx(sample, base, [Layer("x", {"toc": "on"})])
    prs = Presentation(str(out))
    assert "目录" in _texts(prs.slides[1])
    slide = _band_slide(prs)
    assert slide is prs.slides[2]
    texts = _texts(slide)
    seq = [t for k, v in client_band_items(sample) for t in (k, v)]
    pos = [texts.index(t) for t in seq]
    assert pos == sorted(pos)


@pytest.mark.parametrize("scale", SCALES)
@pytest.mark.parametrize("variant", VARIANTS)
def test_ppt_cover_matrix_long_title_subtitle_and_facts_pass_qa(sample, base, scale, variant):
    _crowded(sample)
    for pos in ("top-left", "bottom-right"):
        layers = [Layer("x", {"type.scale": scale, "cover.variant": variant})]
        layers.append(Layer("y", {"brand.logo_position": pos, "brand.org_position": pos}))
        out, style = _pptx(sample, base, layers)
        qa = qa_pptx(out, required_texts(sample), build_theme(style, "pptx"), M)
        assert qa["status"] == "passed", (scale, variant, pos, qa)
        prs = Presentation(str(out))
        assert sample["subtitle"] in _texts(prs.slides[0])
        texts = _texts(_band_slide(prs))
        assert all(f["value"] in texts for f in LONG_FACTS)
        out.unlink()


@pytest.mark.parametrize("scale", SCALES)
def test_band_cells_fit_their_longest_token(sample, scale):
    _crowded(sample)
    style = resolve([Layer("x", {"type.scale": scale})], sample).style
    ctx = make_ctx(sample, build_theme(style, "pptx"), M, __import__("pathlib").Path("."))
    band = InfoBand(tuple(client_band_items(sample)))
    geo = band_geometry(band, BODY_W, ctx)
    t = ctx.theme
    for j, (label, value) in enumerate(band.items):
        c = j % geo.cols
        vw = geo.col_w[c] - geo.label_w[c] - geo.inner
        for tok in nobreak_tokens(value):
            assert M.width(tok, t.body, True) * SAFETY <= vw + 0.01, (tok, vw)
        assert M.lines(label, geo.label_w[c], t.caption) == 1, label


def test_band_puts_label_and_value_on_one_row(sample):
    style = resolve([], sample).style
    ctx = make_ctx(sample, build_theme(style, "pptx"), M, __import__("pathlib").Path("."))
    geo = band_geometry(InfoBand(tuple(client_band_items(sample))), BODY_W, ctx)
    t = ctx.theme
    assert geo.row_h == [pytest.approx(max(lh(t.body), lh(t.caption)))] * len(geo.row_h)


def _LW(s):  # noqa: N802
    return M.width(s, 11) * SAFETY


def _VW(_j, s):  # noqa: N802
    return M.width(s, 14, True) * SAFETY


def test_band_columns_prefers_one_line_values_and_keeps_order():
    items = [("甲", "短"), ("乙", "短"), ("丙", "短"), ("丁", "短"), ("戊", "短")]
    n, widths, labs = band_columns(items, 800.0, 20.0, _LW, _VW, 6.0)
    assert n == 4 and abs(sum(widths) + 3 * 20.0 - 800.0) < 1e-6
    assert labs == [pytest.approx(_LW(k)) for k, _ in items[:4]]
    token = "2026-09-28T07:30:00"
    long = [("甲", "很长的说明文字" * 6), ("乙", f"截至 {token}")]
    n, widths, labs = band_columns(long, 300.0, 20.0, _LW, _VW, 6.0)
    assert n == 2 and widths[1] - labs[1] - 6.0 >= M.width(token, 14) * SAFETY > 110


def test_cover_type_is_capped_at_the_standard_scale(sample, base):
    def title_size(scale):
        out, _ = _pptx(sample, base, [Layer("x", {"type.scale": scale})])
        cover = Presentation(str(out)).slides[0]
        (sh,) = [
            s for s in cover.shapes if s.has_text_frame and s.text_frame.text == sample["title"]
        ]
        size = sh.text_frame.paragraphs[0].runs[0].font.size.pt
        out.unlink()
        return size

    assert title_size("large") == title_size("standard")


def _pdf_html(sample, base, layers=()):
    style = resolve(list(layers), sample).style
    return build_html(sample, style, build_theme(style, "pdf"), base)[0]


def _pdf_cover(html):
    return html[html.index('<section class="cover') : html.index("</section>")]


@pytest.mark.parametrize("variant", VARIANTS)
def test_pdf_cover_holds_exactly_the_cover_elements(sample, base, variant):
    sample["brand"]["logo_path"] = "logo.png"
    cover = _pdf_cover(_pdf_html(sample, base, [Layer("x", {"cover.variant": variant})]))
    text = re.sub(r"<svg.*?</svg>", "", cover, flags=re.S)
    shown = [s for s in re.split(r"<[^>]+>", text) if s.strip()]
    want = [
        sample["brand"]["org_name"],
        sample["title"],
        sample["subtitle"],
        "客户",
        sample["client"]["name"],
        cover_meta_line(sample),
    ]
    assert sorted(shown) == sorted(want)
    assert "<polyline" not in cover and "<path" not in cover


def test_pdf_cover_deco_stays_inside_the_page(sample, base):
    html = _pdf_html(sample, base)
    css = html[html.index("<style>") : html.index("</style>")]
    rule = re.search(r"\.cover \.deco \{([^}]*)\}", css).group(1)
    assert not re.search(r":\s*-", rule), rule
    svg = re.search(r'<div class="deco">(<svg.*?</svg>)', html, re.S).group(1)
    vb = [float(x) for x in re.search(r'viewBox="([^"]+)"', svg).group(1).split()]
    for cx, cy, r in re.findall(r'cx="([\d.]+)" cy="([\d.]+)" r="([\d.]+)"', svg):
        cx, cy, r = float(cx), float(cy), float(r)
        assert cx - r >= vb[0] and cy - r >= vb[1]
        assert cx + r <= vb[0] + vb[2] and cy + r <= vb[1] + vb[3]


def test_pdf_band_follows_toc_and_precedes_the_first_section(sample, base):
    _crowded(sample)
    html = _pdf_html(sample, base, [Layer("x", {"toc": "on"})])
    toc, band = html.index('class="toc"'), html.index('id="client-info"')
    first = html.index(f'id="sec-{sample["sections"][0]["id"]}"')
    assert toc < band < first
    section = html[band:first]
    seq = [t for k, v in client_band_items(sample) for t in (k, v)]
    pos = [section.index(t) for t in seq]
    assert pos == sorted(pos)
    cols = re.search(r'class="band" style="grid-template-columns:([^"]+)"', section).group(1)
    assert len(re.findall(r"[\d.]+pt", cols)) >= 2


def test_band_slack_evens_the_columns_out():
    items = [("甲", "短"), ("乙", "稍长一点的值"), ("丙", "短"), ("丁", "短")]
    n, widths, _ = band_columns(items, 800.0, 20.0, _LW, _VW, 6.0)
    assert n == 4 and max(widths) - min(widths) < 1e-6


def test_first_section_follows_the_band_even_with_little_room(sample):
    from pathlib import Path

    from hpr.ppt_cover import band_section
    from hpr.ppt_layout import BAND_ID, BODY_TOP, lh, paginate
    from hpr.prims import Paragraph, SubHeading

    style = resolve([], sample).style
    ctx = make_ctx(sample, build_theme(style, "pptx"), M, Path("."))
    per = int((ctx.body_bottom - BODY_TOP) // lh(ctx.theme.body))
    avail = ctx.body_bottom - BODY_TOP

    def room_after(k):
        sample["client"]["facts"] = [{"label": f"项目{i}", "value": "数值"} for i in range(k)]
        band = band_section(sample)
        return band, avail - band_geometry(band[0][1][0][0], BODY_W, ctx).height

    band, room = next(room_after(k) for k in range(1, 80) if room_after(k)[1] < 0.3 * avail)
    t = ctx.theme
    assert room > lh(t.heading) + t.gap_xs + t.gap_l + 3 * lh(t.body), room  # heading + 3 lines
    sec = {"id": "s1", "title": "第一节", "blocks": []}
    text = "\n".join(f"第{i}行" for i in range(per))
    pages = paginate([*band, (sec, [(Paragraph(text), "sections[0].blocks[0]")])], ctx)
    assert pages[0].section_id == BAND_ID
    assert any(pl.path == "sections[0]" for pl in pages[0].placed), "section 1 must start here"
    assert [pl.prim for pl in pages[0].placed][1] == SubHeading("第一节")
