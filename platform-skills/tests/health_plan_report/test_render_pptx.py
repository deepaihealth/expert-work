import json
import subprocess
import sys
from pathlib import Path

import pytest
from hpr.content import required_texts
from hpr.measure import Measurer
from hpr.ppt_draw import BODY, Canvas, render_pptx
from hpr.qa import missing_texts, pptx_bounds, pptx_corpus, pptx_overflow, qa_pptx
from hpr.style import Layer, resolve
from hpr.theme import build_theme
from PIL import Image as PILImage
from pptx import Presentation
from pptx.oxml.ns import qn

SKILL = Path(__file__).resolve().parents[2] / "health-plan-report"
RENDER = SKILL / "scripts" / "render.py"
M = Measurer("/nonexistent")


@pytest.fixture
def workdir(tmp_path, sample):
    PILImage.new("RGB", (640, 360), "#DDEEEE").save(tmp_path / "trend-note.png")
    (tmp_path / "plan.json").write_text(json.dumps(sample, ensure_ascii=False), encoding="utf-8")
    return tmp_path


def _render(content, tmp, layers=()):
    style = resolve(list(layers), content).style
    out = tmp / "out.pptx"
    warnings = render_pptx(content, style, out, tmp, M)
    return out, style, warnings


def test_sample_renders_and_passes_qa(workdir, sample):
    out, style, warnings = _render(sample, workdir)
    prs = Presentation(str(out))
    assert prs.slide_width == 12192000 and prs.slide_height == 6858000
    assert len(prs.slides) >= 8
    qa = qa_pptx(out, required_texts(sample), build_theme(style, "pptx"), M)
    assert qa["status"] == "passed", qa
    assert warnings == []


def test_every_required_text_is_found_and_detector_bites(workdir, sample):
    out, _, _ = _render(sample, workdir)
    corpus = pptx_corpus(out)
    assert missing_texts(required_texts(sample), corpus) == []
    assert missing_texts(["一段根本不存在的文字"], corpus) == ["一段根本不存在的文字"]


def test_all_shapes_inside_slide(workdir, sample):
    out, _, _ = _render(sample, workdir)
    assert pptx_bounds(out) == []


def test_bounds_detector_bites(tmp_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = 12192000, 6858000
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    Canvas(slide, build_theme(resolve([]).style, "pptx")).text(
        900, 500, 200, 30, "越界", 14, "#000000"
    )
    p = tmp_path / "b.pptx"
    prs.save(str(p))
    assert pptx_bounds(p)


def test_overflow_detector_bites(tmp_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = 12192000, 6858000
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    Canvas(slide, build_theme(resolve([]).style, "pptx")).text(
        40, 40, 80, 20, "很长的文字" * 20, 14, "#000000"
    )
    p = tmp_path / "o.pptx"
    prs.save(str(p))
    assert pptx_overflow(p, M)


def test_cjk_font_on_every_run(workdir, sample):
    out, _, _ = _render(sample, workdir)
    prs = Presentation(str(out))
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for run in para.runs:
                        ea = run.font._element.find(qn("a:ea"))
                        assert ea is not None and ea.get("typeface") == "微软雅黑"
                        assert run.font.name == "Arial"


def test_body_shapes_are_named(workdir, sample):
    out, _, _ = _render(sample, workdir)
    names = {sh.name for s in Presentation(str(out)).slides for sh in s.shapes}
    assert BODY in names
    assert names <= {
        "hpr:body",
        "hpr:chrome",
        "hpr:deco",
        "hpr:chart",
        "hpr:image",
        "hpr:logo",
        "hpr:media",
    }


@pytest.mark.parametrize("variant", ["band", "split", "minimal"])
@pytest.mark.parametrize("pos", ["top-left", "top-right", "bottom-left", "bottom-right", "center"])
def test_cover_variants_and_logo_positions(workdir, sample, variant, pos):
    PILImage.new("RGB", (300, 120), "#0B4F5C").save(workdir / "logo.png")
    sample["brand"]["logo_path"] = "logo.png"
    layers = [Layer("x", {"cover.variant": variant, "brand.logo_position": pos, "toc": "on"})]
    out, style, warnings = _render(sample, workdir, layers)
    qa = qa_pptx(out, required_texts(sample), build_theme(style, "pptx"), M)
    assert qa["status"] == "passed", qa
    assert warnings == []
    logos = [sh for sh in Presentation(str(out)).slides[0].shapes if sh.name == "hpr:logo"]
    assert len(logos) == 1


def test_missing_logo_is_a_warning_not_a_failure(workdir, sample):
    sample["brand"]["logo_path"] = "nope.png"
    _, _, warnings = _render(sample, workdir)
    assert any("LOGO" in w for w in warnings)


def test_cli_writes_outputs_and_reports(workdir):
    style = workdir / "style.json"
    style.write_text(
        json.dumps({"color.primary": "深蓝", "brand.org_name": "别家"}), encoding="utf-8"
    )
    res = subprocess.run(  # noqa: S603
        [
            sys.executable,
            str(RENDER),
            "--content",
            str(workdir / "plan.json"),
            "--style",
            str(style),
            "--out-dir",
            str(workdir / "out"),
            "--basename",
            "李先生_20260928100000",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert res.returncode == 0, res.stdout + res.stderr
    summary = json.loads(res.stdout)
    assert summary["ok"] is True
    assert Path(summary["files"]["pptx"]).is_file()
    assert [e["key"] for e in summary["not_applied"]] == ["brand.org_name"]
    assert (workdir / "out" / "李先生_20260928100000.qa.json").is_file()
    assert (workdir / "out" / "李先生_20260928100000.params-report.json").is_file()


def test_existing_output_is_not_overwritten(workdir):
    args = [
        sys.executable,
        str(RENDER),
        "--content",
        str(workdir / "plan.json"),
        "--out-dir",
        str(workdir / "out"),
        "--basename",
        "a",
    ]
    assert subprocess.run(args, capture_output=True, check=False).returncode == 0  # noqa: S603
    target = workdir / "out" / "a.pptx"
    before = target.read_bytes()
    again = subprocess.run(args, capture_output=True, text=True, check=False)  # noqa: S603
    assert again.returncode == 1
    assert "已存在" in again.stdout + again.stderr
    assert target.read_bytes() == before


def test_cli_invalid_content_exit_1(workdir, sample):
    del sample["title"]
    bad = workdir / "bad.json"
    bad.write_text(json.dumps(sample, ensure_ascii=False), encoding="utf-8")
    res = subprocess.run(  # noqa: S603
        [
            sys.executable,
            str(RENDER),
            "--content",
            str(bad),
            "--out-dir",
            str(workdir / "o"),
            "--basename",
            "x",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert res.returncode == 1
    assert "title: 缺少必填字段" in json.loads(res.stdout)["errors"]


def test_missing_image_reports_block_path(workdir, sample):
    (workdir / "trend-note.png").unlink()
    (workdir / "plan.json").write_text(json.dumps(sample, ensure_ascii=False), encoding="utf-8")
    res = subprocess.run(  # noqa: S603
        [
            sys.executable,
            str(RENDER),
            "--content",
            str(workdir / "plan.json"),
            "--out-dir",
            str(workdir / "o"),
            "--basename",
            "x",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert res.returncode == 1
    assert any(e.startswith("sections[2].blocks[1]") for e in json.loads(res.stdout)["errors"])


def test_long_title_is_a_clear_error(workdir, sample):
    sample["title"] = "非常长的方案名称" * 12
    with pytest.raises(Exception, match="title"):
        _render(sample, workdir)


def test_basename_with_slash_rejected(workdir):
    res = subprocess.run(  # noqa: S603
        [
            sys.executable,
            str(RENDER),
            "--content",
            str(workdir / "plan.json"),
            "--out-dir",
            str(workdir / "o"),
            "--basename",
            "../x",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert res.returncode == 1


def test_pdf_not_yet_available_message(workdir):
    # T7 实现 PDF 后删除本测试（T7 Step 1 会替换它）  # noqa: RUF003
    res = subprocess.run(  # noqa: S603
        [
            sys.executable,
            str(RENDER),
            "--content",
            str(workdir / "plan.json"),
            "--format",
            "pdf",
            "--out-dir",
            str(workdir / "o"),
            "--basename",
            "x",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert res.returncode == 1
