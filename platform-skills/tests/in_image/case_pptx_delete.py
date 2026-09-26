"""Delete/reorder recipe from pptx/SKILL.md, exec'd verbatim (doc and test cannot drift).

Final review Critical 1: removing a sldId without dropping its relationship leaves the slide
part in the package; after reopen the survivors are renamed and collide with the orphan,
producing duplicate zip entries and silently lost slides.
"""

import collections
import json
import os
import re
import zipfile
from pathlib import Path

import pptx
from _harness import check, run, script

os.chdir("/workspace")
body = (Path(os.environ["EXPERT_WORK_SKILLS_DIR"]) / "pptx" / "SKILL.md").read_text(
    encoding="utf-8"
)
m = re.search(r"```python title=delete-slide\n(.*?)```", body, re.S)
check(m is not None, "SKILL.md has no ```python title=delete-slide block")
snippet = compile(m.group(1), "delete-slide", "exec")


def dup_entries(path: str) -> list[str]:
    names = zipfile.ZipFile(path).namelist()
    return [n for n, c in collections.Counter(names).items() if c > 1]


def titles(path: str) -> list[str]:
    return [s.shapes.title.text for s in pptx.Presentation(path).slides]


prs = pptx.Presentation()
for i in range(1, 6):
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = f"S{i}"
exec(snippet, {"prs": prs})  # noqa: S102 — deletes 3rd slide, moves last to front
prs.save("删后.pptx")
want = ["S5", "S1", "S2", "S4"]
check(dup_entries("删后.pptx") == [], f"dups after delete: {dup_entries('删后.pptx')}")
check(titles("删后.pptx") == want, f"titles={titles('删后.pptx')}")

# rule in SKILL.md: after deleting, save, reopen, then add
again = pptx.Presentation("删后.pptx")
again.slides.add_slide(again.slide_layouts[5]).shapes.title.text = "NEW"
again.save("删后加页.pptx")
check(dup_entries("删后加页.pptx") == [], f"dups after add: {dup_entries('删后加页.pptx')}")
check(titles("删后加页.pptx") == [*want, "NEW"], f"titles={titles('删后加页.pptx')}")

# duplicate_slide.py reopens and re-saves the deleted deck
res = json.loads(
    run(
        [
            "python",
            script("pptx", "duplicate_slide.py"),
            "删后.pptx",
            "删后复制.pptx",
            "--index",
            "1",
        ]
    ).stdout
)
check(res["new_slide_number"] == 2, f"res={res}")
check(dup_entries("删后复制.pptx") == [], f"dups after duplicate: {dup_entries('删后复制.pptx')}")
got = titles("删后复制.pptx")
check(got == ["S5", "S5", "S1", "S2", "S4"], f"titles={got}")
print("PASS case_pptx_delete")
