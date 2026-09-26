import io
import json
import os

import pptx
from _harness import check, run, script
from PIL import Image
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches

os.chdir("/workspace")
p = pptx.Presentation()
for i in range(3):
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = f"第{i + 1}页"
buf = io.BytesIO()
Image.new("RGB", (60, 40), (200, 30, 30)).save(buf, "PNG")
buf.seek(0)
p.slides[1].shapes.add_picture(buf, Inches(1), Inches(2))
p.save("汇报 原件.pptx")
res = json.loads(
    run(
        [
            "python",
            script("pptx", "duplicate_slide.py"),
            "汇报 原件.pptx",
            "汇报 新.pptx",
            "--index",
            "2",
        ]
    ).stdout
)
check(res["new_slide_number"] == 3, f"res={res}")
out = pptx.Presentation("汇报 新.pptx")
titles = [s.shapes.title.text for s in out.slides]
check(titles == ["第1页", "第2页", "第2页", "第3页"], f"titles={titles}")
pics = [sh for sh in out.slides[2].shapes if sh.shape_type == 13]
check(len(pics) == 1 and pics[0].image.blob, "duplicated slide lost its picture")
check(res["shared_parts_warning"] == [], f"picture-only slide should warn nothing: {res}")
out.slides[2].shapes.title.text = "改过的副本"
out.save("汇报 新2.pptx")
check(
    pptx.Presentation("汇报 新2.pptx").slides[1].shapes.title.text == "第2页",
    "editing copy changed original",
)
run(["python", script("pptx", "convert.py"), "汇报 新2.pptx", "--to", "pdf", "--out-dir", "pdf"])
res = json.loads(
    run(
        [
            "python",
            script("pptx", "duplicate_slide.py"),
            "汇报 原件.pptx",
            "汇报 首.pptx",
            "--index",
            "3",
            "--after",
            "0",
        ]
    ).stdout
)
check(
    next(s.shapes.title.text for s in pptx.Presentation("汇报 首.pptx").slides) == "第3页",
    "--after 0 = first",
)
run(
    ["python", script("pptx", "duplicate_slide.py"), "汇报 原件.pptx", "x.pptx", "--index", "9"],
    expect=1,
)

# realistic edge case #1: a slide with a chart — the new slide keeps a working chart, but
# it's the *same* chart-data part as the original (editing one's series edits both), so
# duplicate_slide.py must say so via shared_parts_warning.
p2 = pptx.Presentation()
s2 = p2.slides.add_slide(p2.slide_layouts[6])
chart_data = CategoryChartData()
chart_data.categories = ["A", "B"]
chart_data.add_series("s1", (1, 2))
s2.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0), Inches(0), Inches(4), Inches(3), chart_data
)
p2.save("带图表.pptx")
res = json.loads(
    run(
        [
            "python",
            script("pptx", "duplicate_slide.py"),
            "带图表.pptx",
            "带图表 副本.pptx",
            "--index",
            "1",
        ]
    ).stdout
)
check(res["shared_parts_warning"] == ["图表"], f"chart slide must warn 图表: {res}")
out2 = pptx.Presentation("带图表 副本.pptx")
charts = [sh for sh in out2.slides[1].shapes if sh.has_chart]
check(len(charts) == 1, "duplicated slide lost its chart")

# realistic edge case #2: a slide with speaker notes — a naive implementation drops the
# notes relationship entirely (silent data loss, the brief's own verbatim code), OR relays
# it through the generic shared-rel loop (shares the *same* physical notes part, so a later,
# separate edit to just the copy's notes would silently overwrite the original's too — a
# worse, delayed footgun). The duplicate must get its own independent notes part with the
# same text: no warning (it is not shared), different part identity, and an edit to one
# slide's notes must not leak into the other's after a save/reload round trip.
p3 = pptx.Presentation()
s3 = p3.slides.add_slide(p3.slide_layouts[6])
s3.notes_slide.notes_text_frame.text = "演讲备注：先讲背景"  # noqa: RUF001
p3.save("带备注.pptx")
res = json.loads(
    run(
        [
            "python",
            script("pptx", "duplicate_slide.py"),
            "带备注.pptx",
            "带备注 副本.pptx",
            "--index",
            "1",
        ]
    ).stdout
)
check("备注" not in res["shared_parts_warning"], f"notes must not be shared: {res}")
check(res["notes_format_lost"] is False, f"normal notes copy must not report format lost: {res}")
out3 = pptx.Presentation("带备注 副本.pptx")
notes0, notes1 = out3.slides[0].notes_slide, out3.slides[1].notes_slide
check(notes0.part is not notes1.part, "duplicated slide shares the same notes part")
check(notes0.part.partname != notes1.part.partname, "duplicated slide's notes part not independent")
check(
    notes1.notes_text_frame.text == "演讲备注：先讲背景",  # noqa: RUF001
    f"notes text wrong: {notes1.notes_text_frame.text!r}",
)
# editing the copy's notes and reloading must leave the original's notes untouched.
out3.slides[1].notes_slide.notes_text_frame.text = "改过的副本备注"
out3.save("带备注 副本2.pptx")
out3_reloaded = pptx.Presentation("带备注 副本2.pptx")
check(
    out3_reloaded.slides[0].notes_slide.notes_text_frame.text == "演讲备注：先讲背景",  # noqa: RUF001
    "editing the copy's notes changed the original's notes",
)
check(
    out3_reloaded.slides[1].notes_slide.notes_text_frame.text == "改过的副本备注",
    "edit to the copy's notes did not persist",
)

# realistic edge case #2b: a notes placeholder present but with no <p:txBody> at all (some
# other authoring tool can produce this; python-pptx itself never omits it once a notes
# slide is touched) must not raise — duplicate_slide.py falls back to a plain-text copy and
# discloses that explicitly via notes_format_lost, instead of a broad except silently
# swallowing this (or any unrelated real bug) behind the same "best effort" fallback.
p3b = pptx.Presentation()
s3b = p3b.slides.add_slide(p3b.slide_layouts[6])
notes_ph = s3b.notes_slide.notes_placeholder
notes_ph._element.remove(notes_ph._element.txBody)
p3b.save("备注无正文.pptx")
res3b = json.loads(
    run(
        [
            "python",
            script("pptx", "duplicate_slide.py"),
            "备注无正文.pptx",
            "备注无正文 副本.pptx",
            "--index",
            "1",
        ]
    ).stdout
)
check(res3b["notes_format_lost"] is True, f"missing txBody must report format lost: {res3b}")
out3b = pptx.Presentation("备注无正文 副本.pptx")
new_slide3b = out3b.slides[1]
check(new_slide3b.has_notes_slide, "no-txBody fallback lost the notes slide entirely")
check(
    new_slide3b.notes_slide.notes_text_frame.text == "",
    "no-txBody fallback should copy empty text, not crash or invent content",
)

# realistic edge case #3: an external hyperlink inside a text run must still point to the
# right address on the copy (its r:id is remapped, not left dangling).
p4 = pptx.Presentation()
s4 = p4.slides.add_slide(p4.slide_layouts[6])
tb = s4.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
hrun = tb.text_frame.paragraphs[0].add_run()
hrun.text = "点击访问"
hrun.hyperlink.address = "https://example.com/report"
p4.save("带超链接.pptx")
run(
    [
        "python",
        script("pptx", "duplicate_slide.py"),
        "带超链接.pptx",
        "带超链接 副本.pptx",
        "--index",
        "1",
    ]
)
out4 = pptx.Presentation("带超链接 副本.pptx")
found_link = False
for shp in out4.slides[1].shapes:
    if shp.has_text_frame:
        for para in shp.text_frame.paragraphs:
            for r in para.runs:
                if r.hyperlink.address:
                    found_link = r.hyperlink.address == "https://example.com/report"
check(found_link, "hyperlink address lost or wrong after duplicate")

# realistic edge case #4: a slide edited over time (shapes added and removed) can end up
# with a relationship id that is numerically far from what a fresh, from-scratch count
# would assign — e.g. the picture is "rId5" because three now-deleted relationships used
# up rId2-rId4 first. duplicate_slide.py always builds the new slide from a *fresh* part
# (only rId1, for its own layout link), so the picture's relationship there is "rId2" —
# a real per-attribute r:embed remap is the only thing that keeps the clone connected to
# its own copy of the relationship instead of a stale/nonexistent id from the source.
p5 = pptx.Presentation()
s5 = p5.slides.add_slide(p5.slide_layouts[6])
for i in range(3):
    s5.part.relate_to(f"https://example.com/throwaway{i}", RT.HYPERLINK, is_external=True)
buf5 = io.BytesIO()
Image.new("RGB", (60, 40), (200, 30, 30)).save(buf5, "PNG")
buf5.seek(0)
s5.shapes.add_picture(buf5, Inches(1), Inches(1))  # lands on rId5, not rId2
for i in range(3):
    throwaway_rid = next(
        rid
        for rid, rel in s5.part.rels.items()
        if getattr(rel, "target_ref", None) == f"https://example.com/throwaway{i}"
    )
    s5.part.drop_rel(throwaway_rid)  # leaves the picture's rId5 orphaned-looking but valid
p5.save("id 偏大.pptx")
run(
    [
        "python",
        script("pptx", "duplicate_slide.py"),
        "id 偏大.pptx",
        "id 偏大 副本.pptx",
        "--index",
        "1",
    ]
)
out5 = pptx.Presentation("id 偏大 副本.pptx")
new_slide5 = out5.slides[1]
pics5 = [sh for sh in new_slide5.shapes if sh.shape_type == 13]
check(len(pics5) == 1, "high-id slide lost its picture")
check(
    len(pics5[0].image.blob) > 0,
    "high-id: picture r:embed not remapped to the new slide's own relationship id",
)

# realistic edge case #5: duplicating the same source slide twice in a row (once into an
# intermediate file, then duplicating a slide of *that* file again) must not corrupt
# relationship ids or crash — each invocation is a fresh process/file.
res = json.loads(
    run(
        [
            "python",
            script("pptx", "duplicate_slide.py"),
            "带图表.pptx",
            "带图表 副本.pptx",
            "--index",
            "1",
        ]
    ).stdout
)
res2 = json.loads(
    run(
        [
            "python",
            script("pptx", "duplicate_slide.py"),
            "带图表 副本.pptx",
            "带图表 副本2.pptx",
            "--index",
            "1",
        ]
    ).stdout
)
check(res2["shared_parts_warning"] == ["图表"], f"second duplicate must still warn: {res2}")
out6 = pptx.Presentation("带图表 副本2.pptx")
check(len(out6.slides) == 3, f"expected 3 slides after duplicating twice, got {len(out6.slides)}")
# slides[1] is the actual product of the *second* duplicate_slide.py call (--index 1 with no
# --after inserts right after position 1); slides[2] is the untouched first-generation copy
# carried over from before this call, so asserting on it wouldn't catch a second-call bug.
check(
    len([sh for sh in out6.slides[1].shapes if sh.has_chart]) == 1,
    "second duplicate's own output slide lost its chart",
)

print("PASS case_pptx_duplicate")
