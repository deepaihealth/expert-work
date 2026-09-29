"""Page map in the real image: every section's first listed page shows its title, the ranges
are contiguous and together cover the whole document (PDF anchors from real weasyprint)."""

import json
import os
import re
import shutil
from itertools import pairwise
from pathlib import Path

import pdfplumber
from _harness import check, run, script
from PIL import Image
from pptx import Presentation

os.chdir("/workspace")
skill = Path(os.environ["EXPERT_WORK_SKILLS_DIR"]) / "health-plan-report"
shutil.copy(skill / "sample" / "sample-plan.json", "plan.json")
Image.new("RGB", (640, 360), "#DDEEEE").save("trend-note.png")


def squash(s: str) -> str:
    return re.sub(r"\s+", "", s)


def pdf_pages(path: str) -> list[str]:
    with pdfplumber.open(path) as pdf:
        return [squash(p.extract_text() or "") for p in pdf.pages]


def pptx_pages(path: str) -> list[str]:
    out = []
    for slide in Presentation(path).slides:
        parts = []
        for sh in slide.shapes:
            if sh.has_text_frame:
                parts.append(sh.text_frame.text)
            elif getattr(sh, "has_table", False) and sh.has_table:
                parts += [c.text for r in sh.table.rows for c in r.cells]
        out.append(squash("".join(parts)))
    return out


def verify(entries: list[dict], texts: list[str], label: str) -> None:
    covered = set()
    for e in entries:
        pages = e["pages"]
        check(pages == list(range(pages[0], pages[-1] + 1)), f"{label} {e['title']} contiguous")
        check(1 <= pages[0] and pages[-1] <= len(texts), f"{label} {e['title']} in range")
        covered.update(pages)
        if "section" in e:
            check(squash(e["title"]) in texts[pages[0] - 1], f"{label} {e['title']} on first page")
    check(covered == set(range(1, len(texts) + 1)), f"{label} covers every page")


for scale in ("standard", "large"):
    name = f"pm-{scale}"
    Path(f"{name}.style.json").write_text(json.dumps({"type.scale": scale}), "utf-8")
    cmd = ["python", script("health-plan-report", "render.py"), "--content", "plan.json"]
    cmd += ["--style", f"{name}.style.json", "--format", "both", "--out-dir", "out"]
    res = json.loads(run([*cmd, "--basename", name], expect=0, timeout=600).stdout)
    check(res["ok"], f"{scale} renders")
    verify(res["page_map"]["pptx"], pptx_pages(res["files"]["pptx"]), f"pptx/{scale}")
    verify(res["page_map"]["pdf"], pdf_pages(res["files"]["pdf"]), f"pdf/{scale}")
    sections = [e for e in res["page_map"]["pdf"] if "section" in e]
    for a, b in pairwise(sections):
        check(a["pages"][-1] <= b["pages"][0], f"pdf/{scale} {a['title']} ends before next starts")
    with pdfplumber.open(res["files"]["pdf"]) as pdf:
        toc = pdf.pages[1].extract_text() or ""
    for i, e in enumerate(sections, start=1):
        m = re.search(rf"{i:02d}\s*{re.escape(e['title'])}\.*\s*(\d+)", toc)
        check(bool(m) and int(m.group(1)) == e["pages"][0], f"pdf/{scale} TOC page of {e['title']}")
    multi = [e for e in res["page_map"]["pdf"] if len(e["pages"]) > 1]
    check(bool(multi), f"pdf/{scale} has a section spanning pages (end markers exercised)")
print("PASS case_hpr_page_map")
