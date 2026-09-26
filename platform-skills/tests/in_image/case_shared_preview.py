import json
import os
from pathlib import Path

import docx
from _harness import check, png_is_not_blank, run, script

os.chdir("/workspace")
d = docx.Document()
for i in range(5):
    d.add_heading(f"第 {i + 1} 页 标题", 1)
    d.add_paragraph("正文内容 " * 50)
    if i < 4:
        d.add_page_break()
d.save("长 文档.docx")

res = json.loads(run(["python", script("docx", "preview.py"), "长 文档.docx"]).stdout)
check(res["pages_total"] == 5, f"pages_total={res['pages_total']}")
check(len(res["images"]) == 3, f"default should render 3 pages, got {len(res['images'])}")
for img in res["images"]:
    check(Path(img).is_file() and png_is_not_blank(Path(img)), f"blank or missing {img}")
check(
    Path(res["images"][0]).resolve().parent == Path("/workspace/长 文档_preview"),
    f"default out dir wrong: {res['images'][0]}",
)

res = json.loads(
    run(
        [
            "python",
            script("docx", "preview.py"),
            "长 文档.docx",
            "--pages",
            "all",
            "--out-dir",
            "a/b/c",
        ]
    ).stdout
)
check(len(res["images"]) == 5, "all pages")
res = json.loads(
    run(
        [
            "python",
            script("docx", "preview.py"),
            "长 文档.docx",
            "--pages",
            "2,4",
            "--out-dir",
            "sel",
        ]
    ).stdout
)
check(len(res["images"]) == 2, "selected pages")
run(["python", script("docx", "preview.py"), "不存在.docx"], expect=1)
print("PASS case_shared_preview")
