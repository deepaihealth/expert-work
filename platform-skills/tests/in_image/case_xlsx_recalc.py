import json
import os
import re
import zipfile

import openpyxl
from _harness import check, run, script

os.chdir("/workspace")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "数据"
ws["A1"], ws["A2"], ws["A3"] = 2, 3, "=A1*A2"
ws2 = wb.create_sheet("错误")
ws2["B2"] = "=1/0"
ws2["B3"] = "=数据!A3+1"
wb.save("预算 表.xlsx")
res = json.loads(
    run(
        ["python", script("xlsx", "recalc.py"), "预算 表.xlsx", "重算/预算 表.xlsx"], expect=2
    ).stdout
)
check(res["formulas"] == 3, f"formulas={res['formulas']}")
check(
    res["errors"] == [{"sheet": "错误", "cell": "B2", "value": "#DIV/0!"}],
    f"errors={res['errors']}",
)
vals = openpyxl.load_workbook("重算/预算 表.xlsx", data_only=True)
check(
    vals["数据"]["A3"].value == 6 and vals["错误"]["B3"].value == 7,
    "cached values missing after recalc",
)
check(openpyxl.load_workbook("重算/预算 表.xlsx")["数据"]["A3"].value == "=A1*A2", "formula lost")

# extra: exercise ALL 7 _ERRORS strings through one real recalc.py run (the test above only
# hits #DIV/0!), one formula per error, asserting each is reported against the right cell.
wb5 = openpyxl.Workbook()
ws5 = wb5.active
ws5.title = "错误全集"
ws5["A1"] = "=1/0"  # #DIV/0!
ws5["A2"] = "=INDEX(A1:B1,1,5)"  # #REF! (column index out of the 2-column range)
ws5["A3"] = '="a"+1'  # #VALUE!
ws5["A4"] = "=NOPE()"  # #NAME? (undefined function)
ws5["A5"] = "=NA()"  # #N/A
ws5["A6"] = "=ASIN(2)"  # #NUM! (out of domain)
ws5["A7"] = "=SUM(A1 B1)"  # #NULL! (space = intersection of non-overlapping cells)
wb5.save("全部错误.xlsx")
res_all = json.loads(
    run(
        ["python", script("xlsx", "recalc.py"), "全部错误.xlsx", "全部错误_出.xlsx"], expect=2
    ).stdout
)
expected_all = {
    "A1": "#DIV/0!",
    "A2": "#REF!",
    "A3": "#VALUE!",
    "A4": "#NAME?",
    "A5": "#N/A",
    "A6": "#NUM!",
    "A7": "#NULL!",
}
got_all = {e["cell"]: e["value"] for e in res_all["errors"] if e["sheet"] == "错误全集"}
check(got_all == expected_all, f"not all 7 _ERRORS reproduced: {got_all}")

wb2 = openpyxl.Workbook()
wb2.active["A1"] = "=1+1"
wb2.save("ok.xlsx")
run(["python", script("xlsx", "recalc.py"), "ok.xlsx", "ok2.xlsx"], expect=0)
run(["python", script("xlsx", "recalc.py"), "ok.xlsx", "ok.xlsx"], expect=1)

# extra: a workbook with no formulas at all must report formulas=0, errors=[], exit 0
# (recalc.py's counting loop has nothing else to exercise this path).
wb3 = openpyxl.Workbook()
wb3.active["A1"] = "纯文本"
wb3.active["A2"] = 42
wb3.save("无公式.xlsx")
res_noformula = json.loads(
    run(["python", script("xlsx", "recalc.py"), "无公式.xlsx", "无公式_出.xlsx"]).stdout
)
check(res_noformula == {"formulas": 0, "errors": []}, f"no-formula file: {res_noformula}")

# extra: the OOXMLRecalcMode xcu (_office.py's recalc=True branch) must override a
# genuinely STALE cached value, not just fill in ones that were never there (which
# LibreOffice does unconditionally -- see recalc.py's module docstring for the full
# investigation). Get a workbook with a REAL cached value via one recalc.py pass, hand-
# forge that cached <v> to a wrong number without touching its <f>, then confirm the next
# recalc.py pass corrects it back to the true recomputed value.
wb4 = openpyxl.Workbook()
ws4 = wb4.active
ws4.title = "数据"
ws4["A1"], ws4["A2"], ws4["A3"] = 2, 3, "=A1*A2"
wb4.save("新鲜.xlsx")
run(["python", script("xlsx", "recalc.py"), "新鲜.xlsx", "带缓存.xlsx"])

with zipfile.ZipFile("带缓存.xlsx") as z:
    names = z.namelist()
    parts = {n: z.read(n) for n in names}
sheet_name = next(n for n in names if n.startswith("xl/worksheets/sheet") and n.endswith(".xml"))
xml = parts[sheet_name].decode("utf-8")
patched, count = re.subn(r'(<c r="A3"[^>]*><f[^>]*>[^<]*</f>)<v>[^<]*</v>', r"\g<1><v>999</v>", xml)
check(count == 1, "regex did not find exactly one A3 cached value to patch")
parts[sheet_name] = patched.encode("utf-8")
with zipfile.ZipFile("陈旧.xlsx", "w", zipfile.ZIP_DEFLATED) as zf:
    for n in names:
        zf.writestr(n, parts[n])
check(
    openpyxl.load_workbook("陈旧.xlsx", data_only=True)["数据"]["A3"].value == 999,
    "hand-forged stale cache did not take",
)
run(["python", script("xlsx", "recalc.py"), "陈旧.xlsx", "修正.xlsx"])
fixed = openpyxl.load_workbook("修正.xlsx", data_only=True)["数据"]["A3"].value
check(fixed == 6, "recalc.py did not override the stale cached value")

print("PASS case_xlsx_recalc")
