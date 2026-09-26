"""Merged table cells are visited once (final review Important 2).

``row.cells`` yields a horizontally merged cell once per spanned column and a vertically
merged cell once per spanned row; a replacement containing its own search text then
stacks (甲公司 -> 甲某某某某某某公司) and counts are inflated.
"""

import json
import os

import docx
from _harness import check, run, script

os.chdir("/workspace")
d = docx.Document()
t = d.add_table(rows=3, cols=3)
t.cell(0, 0).merge(t.cell(0, 2))  # horizontal: 3 grid columns
t.cell(0, 0).text = "甲公司"
t.cell(1, 0).merge(t.cell(2, 0))  # vertical: 2 rows
t.cell(1, 0).text = "乙公司"
t.cell(1, 1).text = "丙公司"
inner = t.cell(2, 2).add_table(rows=1, cols=2)
inner.cell(0, 0).merge(inner.cell(0, 1))  # nested, horizontal
inner.cell(0, 0).text = "丁公司"
d.save("合并.docx")
with open("merged_rules.json", "w", encoding="utf-8") as f:
    json.dump([{"find": "公司", "replace": "某某公司"}], f, ensure_ascii=False)
res = json.loads(
    run(
        [
            "python",
            script("docx", "replace_text.py"),
            "合并.docx",
            "合并_出.docx",
            "--rules",
            "merged_rules.json",
        ]
    ).stdout
)
check(res["rules"][0]["count"] == 4, f"count should be 4 (one per real cell): {res}")
out = docx.Document("合并_出.docx").tables[0]
check(out.cell(0, 0).text == "甲某某公司", f"horizontal merge: {out.cell(0, 0).text!r}")
check(out.cell(1, 0).text == "乙某某公司", f"vertical merge: {out.cell(1, 0).text!r}")
check(out.cell(1, 1).text == "丙某某公司", f"plain cell: {out.cell(1, 1).text!r}")
nested = out.cell(2, 2).tables[0].cell(0, 0).text
check(nested == "丁某某公司", f"nested merge: {nested!r}")

# fill_template shares the same iterator
d2 = docx.Document()
t2 = d2.add_table(rows=2, cols=3)
t2.cell(0, 0).merge(t2.cell(0, 2))
t2.cell(0, 0).text = "甲方：{{甲方}}"  # noqa: RUF001
t2.cell(1, 0).text = "{{日期}}"
d2.save("合并模板.docx")
with open("merged_data.json", "w", encoding="utf-8") as f:
    json.dump({"甲方": "{{甲方}}有限公司", "日期": "2026-09-28"}, f, ensure_ascii=False)
res = json.loads(
    run(
        [
            "python",
            script("docx", "fill_template.py"),
            "合并模板.docx",
            "合并成品.docx",
            "--data",
            "merged_data.json",
        ]
    ).stdout
)
check(sorted(res["filled"]) == sorted(["日期", "甲方"]), f"{res}")
out2 = docx.Document("合并成品.docx").tables[0]
got = out2.cell(0, 0).text
check(got == "甲方：{{甲方}}有限公司", f"merged fill stacked: {got!r}")  # noqa: RUF001
check(out2.cell(1, 0).text == "2026-09-28", f"date: {out2.cell(1, 0).text!r}")
print("PASS case_docx_merged")
