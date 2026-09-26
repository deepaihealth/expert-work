import json
import os

import docx
from _harness import check, run, script

os.chdir("/workspace")
d = docx.Document()
p = d.add_paragraph()
r1 = p.add_run("客户")
r1.bold = True
p.add_run("姓名：")  # noqa: RUF001
p.add_run("张三")
t = d.add_table(rows=1, cols=1)
t.cell(0, 0).paragraphs[0].add_run("客户姓名：张三")  # noqa: RUF001
inner = t.cell(0, 0).add_table(rows=1, cols=1)
inner.cell(0, 0).paragraphs[0].add_run("嵌套 张三")
d.sections[0].header.paragraphs[0].add_run("页眉 张三")
d.save("原 件.docx")
with open("rules.json", "w", encoding="utf-8") as f:
    json.dump(
        [
            {"find": "客户姓名：张三", "replace": "客户姓名：李四"},  # noqa: RUF001
            {"find": "张三", "replace": "李四"},
            {"find": "不存在的词", "replace": "x"},
        ],
        f,
        ensure_ascii=False,
    )
res = json.loads(
    run(
        [
            "python",
            script("docx", "replace_text.py"),
            "原 件.docx",
            "输出/新 件.docx",
            "--rules",
            "rules.json",
        ]
    ).stdout
)
check(res["unmatched"] == ["不存在的词"], f"unmatched={res['unmatched']}")
out = docx.Document("输出/新 件.docx")
first = out.paragraphs[0]
check(first.text == "客户姓名：李四", f"split-run replace wrong: {first.text!r}")  # noqa: RUF001
check(first.runs[0].bold is True, "first run lost bold")
cell = out.tables[0].cell(0, 0)
check(cell.paragraphs[0].text == "客户姓名：李四", f"table text {cell.paragraphs[0].text!r}")  # noqa: RUF001
check(cell.tables[0].cell(0, 0).paragraphs[0].text == "嵌套 李四", "nested table")
check(out.sections[0].header.paragraphs[0].text == "页眉 李四", "header")
check(docx.Document("原 件.docx").paragraphs[0].text == "客户姓名：张三", "original was modified")  # noqa: RUF001
run(
    [
        "python",
        script("docx", "replace_text.py"),
        "原 件.docx",
        "原 件.docx",
        "--rules",
        "rules.json",
    ],
    expect=1,
)
print("PASS case_docx_replace")
