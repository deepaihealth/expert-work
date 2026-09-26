import json
import os
import threading
from pathlib import Path

import docx
import openpyxl
import pptx
from _harness import check, run, script

os.chdir("/workspace")
d = docx.Document()
d.add_paragraph("中文段落 测试")
d.save("张三 方案.docx")
p = pptx.Presentation()
s = p.slides.add_slide(p.slide_layouts[1])
s.shapes.title.text = "中文标题"
p.save("汇报 一.pptx")
wb = openpyxl.Workbook()
wb.active["A1"] = "中文"
wb.save("数据 表.xlsx")

for name in ("张三 方案.docx", "汇报 一.pptx", "数据 表.xlsx"):
    out = json.loads(
        run(
            ["python", script("docx", "convert.py"), name, "--to", "pdf", "--out-dir", "out/pdf"]
        ).stdout
    )
    pdf = Path(out["output"])
    check(pdf.is_file() and pdf.stat().st_size > 1000, f"no pdf for {name}")
    check(
        pdf.resolve().is_relative_to(Path("/workspace/out/pdf")),
        f"pdf not under /workspace/out/pdf: {pdf}",
    )

results = []


def _one(src: str, sub: str) -> None:
    results.append(
        run(
            ["python", script("docx", "convert.py"), src, "--to", "pdf", "--out-dir", sub]
        ).returncode
    )


ts = [threading.Thread(target=_one, args=("张三 方案.docx", f"par{i}")) for i in range(2)]
[t.start() for t in ts]
[t.join() for t in ts]
check(results == [0, 0], f"concurrent conversions: {results}")
run(
    ["python", script("docx", "convert.py"), "张三 方案.docx", "--to", "docx", "--out-dir", "."],
    expect=1,
)
print("PASS case_shared_convert")
