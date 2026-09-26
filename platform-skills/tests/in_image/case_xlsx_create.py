import json
import os
import re
from pathlib import Path

import openpyxl
from _harness import check, png_is_not_blank, run, script

os.chdir("/workspace")
body = (Path(os.environ["EXPERT_WORK_SKILLS_DIR"]) / "xlsx" / "SKILL.md").read_text(
    encoding="utf-8"
)
m = re.search(r"```python title=skeleton\n(.*?)```", body, re.S)
check(m is not None, "SKILL.md has no ```python title=skeleton block")
exec(compile(m.group(1), "skeleton", "exec"), {})  # noqa: S102
out = Path("骨架 示例.xlsx")
check(out.is_file(), "skeleton did not produce 骨架 示例.xlsx")

wb = openpyxl.load_workbook(out)
ws = wb.active
check(ws.freeze_panes == "A2", f"freeze_panes={ws.freeze_panes!r}")
formula_cell = ws["B5"]
check(
    isinstance(formula_cell.value, str) and formula_cell.value.startswith("="),
    f"SUM formula cell wrong: {formula_cell.value!r}",
)
check(ws["A1"].font.bold is True, "header font is not bold")

recalc = json.loads(
    run(["python", script("xlsx", "recalc.py"), str(out), "重算/骨架.xlsx"], expect=0).stdout
)
check(recalc["errors"] == [], f"skeleton produced error cells: {recalc['errors']}")
vals = openpyxl.load_workbook("重算/骨架.xlsx", data_only=True)
check(vals.active["B5"].value == 30000, f"SUM cached value wrong: {vals.active['B5'].value}")

res = json.loads(run(["python", script("xlsx", "preview.py"), str(out)]).stdout)
check(res["images"], "preview produced nothing")
check(png_is_not_blank(Path(res["images"][0])), "preview image is blank")

print("PASS case_xlsx_create")
