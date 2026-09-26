import json
import os

import docx
from _harness import check, run, script

os.chdir("/workspace")
d = docx.Document()
p = d.add_paragraph()
p.add_run("尊敬的 {{客户")
p.add_run("名}}，您好")  # noqa: RUF001
d.add_paragraph("机构：{{机构名}}  日期：{{日期}}")  # noqa: RUF001
# spaced placeholder ("{{ key }}"), same key as "客户名" above
d.add_paragraph("附言：{{ 客户名 }} 敬上")  # noqa: RUF001
d.save("模板.docx")
res = json.loads(run(["python", script("docx", "fill_template.py"), "模板.docx", "--list"]).stdout)
check(sorted(res["placeholders"]) == ["客户名", "日期", "机构名"], f"placeholders={res}")
with open("data.json", "w", encoding="utf-8") as f:
    json.dump({"客户名": "李四", "机构名": "深护智康", "多余": 1}, f, ensure_ascii=False)
res = json.loads(
    run(
        [
            "python",
            script("docx", "fill_template.py"),
            "模板.docx",
            "成品.docx",
            "--data",
            "data.json",
        ]
    ).stdout
)
check(res["missing"] == ["日期"] and res["unused"] == ["多余"], f"report={res}")
text = "\n".join(pp.text for pp in docx.Document("成品.docx").paragraphs)
check("尊敬的 李四，您好" in text and "机构：深护智康" in text and "{{日期}}" in text, text)  # noqa: RUF001
check("附言：李四 敬上" in text, f"spaced placeholder not filled: {text!r}")  # noqa: RUF001

# null/list/dict/bool values are all rejected (exit 1); only str/int/float are accepted.
with open("data_null.json", "w", encoding="utf-8") as f:
    json.dump({"客户名": None}, f, ensure_ascii=False)
run(
    [
        "python",
        script("docx", "fill_template.py"),
        "模板.docx",
        "废.docx",
        "--data",
        "data_null.json",
    ],
    expect=1,
)
print("PASS case_docx_template")
