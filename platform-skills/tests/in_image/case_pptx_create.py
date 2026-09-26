import json
import os
import re
from pathlib import Path

import pdfplumber
import pptx
from _harness import check, run, script
from pptx.oxml.ns import qn

os.chdir("/workspace")
body = (Path(os.environ["EXPERT_WORK_SKILLS_DIR"]) / "pptx" / "SKILL.md").read_text(
    encoding="utf-8"
)
m = re.search(r"```python title=skeleton\n(.*?)```", body, re.S)
check(m is not None, "SKILL.md has no ```python title=skeleton block")
exec(compile(m.group(1), "skeleton", "exec"), {})  # noqa: S102
out = Path("骨架 示例.pptx")
check(out.is_file(), "skeleton did not produce 骨架 示例.pptx")
prs = pptx.Presentation(str(out))
check(
    prs.slide_width == 12192000 and prs.slide_height == 6858000,
    f"size={prs.slide_width}x{prs.slide_height} (want 16:9 12192000x6858000)",
)
check(len(prs.slides) >= 3, f"skeleton must have >= 3 slides, got {len(prs.slides)}")
found_ea = any(
    run.font._element.find(qn("a:ea")) is not None
    and run.font._element.find(qn("a:ea")).get("typeface") == "微软雅黑"
    for slide in prs.slides
    for shape in slide.shapes
    if shape.has_text_frame
    for para in shape.text_frame.paragraphs
    for run in para.runs
)
check(found_ea, "no run has a:ea typeface=微软雅黑")
has_table = any(shape.has_table for slide in prs.slides for shape in slide.shapes)
has_picture = any(shape.shape_type == 13 for slide in prs.slides for shape in slide.shapes)
check(has_table, "skeleton must show a table")
check(has_picture, "skeleton must show a picture")
pdf = json.loads(
    run(
        ["python", script("pptx", "convert.py"), str(out), "--to", "pdf", "--out-dir", "pdf"]
    ).stdout
)["output"]
with pdfplumber.open(pdf) as pp:
    text = "".join(pg.extract_text() or "" for pg in pp.pages)
check(
    "示例汇报" in text or "汇报" in text,
    f"CJK text not extractable after 微软雅黑 fallback: {text[:200]!r}",
)
res = json.loads(run(["python", script("pptx", "preview.py"), str(out)]).stdout)
check(res["images"], "preview produced nothing")
print("PASS case_pptx_create")
