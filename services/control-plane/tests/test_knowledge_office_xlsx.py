"""B-206 K-18 —— a spreadsheet's Markdown shows what its reader sees.

MarkItDown read ``.xlsx`` with pandas: a merged category was left on its
first row only and the rows below read ``NaN``, an empty cell read ``NaN``,
a percentage read as its stored fraction, a date at midnight read
``00:00:00``, a ``|`` in a cell split the column, a title row became the
header (``Unnamed: 1`` …). Here:

* a cell merged down the rows is copied into every row (across the columns
  it stays in its first cell);
* an empty cell is empty; a row (or trailing column) with nothing is dropped;
* numbers as displayed: percentages, fixed decimals, whole floats;
  dates without a midnight time;
* ``|`` escaped, line breaks folded into spaces;
* a first row holding a single value over a wider table is the sheet's
  title — a heading, not the header;
* a file openpyxl cannot read still goes to MarkItDown.

Made-up data only.
"""

from __future__ import annotations

import datetime
import io

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from control_plane.knowledge import parsing
from control_plane.knowledge.office_xlsx import xlsx_markdown
from control_plane.knowledge.parsing import parse_document


def _book(*sheets: tuple[str, list[list[object]]]) -> tuple[Workbook, list[Worksheet]]:
    book = Workbook()
    book.remove(book.active)
    made = []
    for title, rows in sheets:
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
        made.append(sheet)
    return book, made


def _bytes(book: Workbook) -> bytes:
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_a_plain_sheet_is_a_markdown_table_under_its_name() -> None:
    book, _ = _book(("Stock", [["Item", "Count"], ["Gloves", 12], ["Masks", 30]]))
    assert xlsx_markdown(_bytes(book)) == (
        "## Stock\n| Item | Count |\n| --- | --- |\n| Gloves | 12 |\n| Masks | 30 |"
    )


def test_each_sheet_follows_the_last() -> None:
    book, _ = _book(("One", [["a", "b"], [1, 2]]), ("Two", [["c"], [3]]))
    assert xlsx_markdown(_bytes(book)) == (
        "## One\n| a | b |\n| --- | --- |\n| 1 | 2 |\n\n## Two\n| c |\n| --- |\n| 3 |"
    )


def test_a_category_merged_down_the_rows_is_in_every_row() -> None:
    book, (sheet,) = _book(
        (
            "Plans",
            [
                ["Group", "Plan", "Fee"],
                ["Basic", "A", 10],
                [None, "B", 20],
                [None, "C", 30],
                ["Plus", "D", 40],
            ],
        )
    )
    sheet.merge_cells("A2:A4")
    assert xlsx_markdown(_bytes(book)).splitlines()[1:] == [
        "| Group | Plan | Fee |",
        "| --- | --- | --- |",
        "| Basic | A | 10 |",
        "| Basic | B | 20 |",
        "| Basic | C | 30 |",
        "| Plus | D | 40 |",
    ]


def test_a_merge_across_columns_stays_in_its_first_cell() -> None:
    book, (sheet,) = _book(
        ("Sheet", [["Item", "Spring", "Autumn"], ["Rice", "both seasons", None]])
    )
    sheet.merge_cells("B2:C2")
    assert xlsx_markdown(_bytes(book)).splitlines()[-1] == "| Rice | both seasons |  |"


def test_a_block_merge_is_copied_down_its_first_column_only() -> None:
    book, (sheet,) = _book(
        (
            "Sheet",
            [["Area", "Note", "Code"], ["North", None, "N1"], [None, None, "N2"]],
        )
    )
    sheet.merge_cells("A2:B3")
    assert xlsx_markdown(_bytes(book)).splitlines()[3:] == [
        "| North |  | N1 |",
        "| North |  | N2 |",
    ]


def test_empty_cells_are_empty_and_empty_rows_and_trailing_columns_are_dropped() -> None:
    book, (sheet,) = _book(
        ("Sheet", [["Name", "Phone", None], ["Ann", None, None], [None, None, None], ["Bo", 5]])
    )
    sheet["D9"].number_format = "0.00"  # formatted, never filled: no column, no row
    markdown = xlsx_markdown(_bytes(book))
    assert "NaN" not in markdown and "Unnamed" not in markdown
    assert markdown.splitlines()[1:] == [
        "| Name | Phone |",
        "| --- | --- |",
        "| Ann |  |",
        "| Bo | 5 |",
    ]


@pytest.mark.parametrize(
    ("value", "number_format", "shown"),
    [
        (0.6112, "0.00%", "61.12%"),
        (0.25, "0%", "25%"),
        (0.125, "0.0%", "12.5%"),
        (3.14159, "0.00_ ", "3.14"),
        (2.5, "0.0", "2.5"),
        (12.0, "General", "12"),
        (0.1 + 0.2, "General", "0.3"),
        (1234567, "0_ ", "1234567"),
        (True, "General", "TRUE"),
        (datetime.datetime(2026, 3, 1), "yyyy/mm/dd", "2026-03-01"),
        (datetime.datetime(2026, 3, 1, 9, 30), "yyyy/mm/dd hh:mm", "2026-03-01 09:30"),
        (datetime.datetime(2026, 3, 1, 9, 30, 15), "yyyy-mm-dd hh:mm:ss", "2026-03-01 09:30:15"),
        (datetime.time(8, 5), "hh:mm", "08:05"),
    ],
)
def test_values_read_as_displayed(value: object, number_format: str, shown: str) -> None:
    book, (sheet,) = _book(("Sheet", [["Value"], [value]]))
    sheet["A2"].number_format = number_format
    assert xlsx_markdown(_bytes(book)).splitlines()[-1] == f"| {shown} |"


def test_pipes_are_escaped_and_line_breaks_folded() -> None:
    book, _ = _book(("Sheet", [["Name", "Note"], ["Beta | B", "line one\nline two\r\nthree"]]))
    assert xlsx_markdown(_bytes(book)).splitlines()[-1] == (
        "| Beta \\| B | line one line two three |"
    )


def test_a_single_value_over_a_wider_table_is_the_sheets_title() -> None:
    book, (sheet,) = _book(
        ("Prices", [["Price list 2026", None, None], ["Name", "Unit", "Price"], ["Rice", "kg", 6]])
    )
    sheet.merge_cells("A1:C1")
    assert xlsx_markdown(_bytes(book)) == (
        "## Prices\n### Price list 2026\n| Name | Unit | Price |\n| --- | --- | --- |\n"
        "| Rice | kg | 6 |"
    )


def test_a_one_column_table_keeps_its_first_row_as_the_header() -> None:
    book, _ = _book(("Sheet", [["Names"], ["Ann"], ["Bo"]]))
    assert xlsx_markdown(_bytes(book)).splitlines()[1] == "| Names |"


def test_a_title_alone_is_still_the_sheet() -> None:
    book, _ = _book(("Sheet", [["Only a note", None]]), ("Empty", []))
    assert xlsx_markdown(_bytes(book)) == "## Sheet\n| Only a note |\n| --- |"


def test_parse_document_uses_it_for_xlsx_and_falls_back_to_markitdown(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    book, (sheet,) = _book(("Sheet", [["Group", "Item"], ["Fruit", "Apple"], [None, "Pear"]]))
    sheet.merge_cells("A2:A3")
    raw = _bytes(book)
    assert parse_document("plans.xlsx", raw).splitlines()[-1] == "| Fruit | Pear |"

    def broken(raw: bytes) -> str:
        raise ValueError("unreadable")

    monkeypatch.setattr(parsing, "xlsx_markdown", broken)
    # MarkItDown's own reading, as before K-18.
    assert parse_document("plans.xlsx", raw).splitlines()[-1] == "| NaN | Pear |"
