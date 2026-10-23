"""The per-document cap on parsed text and the server's budget for it
(``control_plane/knowledge/parsed_size.py``).

A compressed upload can unpack to far more text than its size suggests, and
every character costs the server memory after parsing (chunks, a vector per
chunk, the rows written). Over the cap a document fails with
``text_too_large`` — refused in the parse worker for every format, so the text
never reaches the server — and is never retried; the budget keeps concurrent
ingests from adding up past what one document at the cap may use.

Documents are synthetic; the parse pool is the real (spawned) one.
"""

from __future__ import annotations

import asyncio
import io
import json
import pickle
import zipfile
from collections.abc import AsyncIterator, Callable, Iterator, Sequence
from contextlib import asynccontextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4
from xml.sax.saxutils import escape

import pymupdf
import pytest
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from control_plane.knowledge import ingestion, parsed_size
from control_plane.knowledge.ingestion import (
    KnowledgeIngestionRunner,
    failure_code,
    ingest_document_bytes,
    parse_document_in_pool,
    shutdown_parse_pool,
)
from control_plane.knowledge.parsed_size import (
    DEFAULT_MAX_PARSED_CHARS,
    DEFAULT_RAW_BYTES_BUDGET,
    DEFAULT_UPLOAD_MAX_BYTES,
    TOO_LARGE_PATTERN,
    DocumentTooLargeError,
    admitted,
    configure_parsed_size,
    configure_raw_bytes,
    ensure_within_cap,
    parse_within_cap,
    reserve_raw_bytes,
    too_large_message,
)
from control_plane.knowledge.parsing import parse_document
from control_plane.knowledge.recognition import recognizer as recognizer_module
from control_plane.knowledge.recognition.clients import Reading, RecognitionClient
from control_plane.knowledge.recognition.figure_extract import parse_office_figures
from control_plane.knowledge.recognition.figure_pipeline import FigureDescriber
from control_plane.knowledge.recognition.figure_storage import FigureStorage
from control_plane.knowledge.recognition.recognizer import (
    DocumentRecognizer,
    ModelChoice,
    RecognitionSettings,
)
from control_plane.knowledge.recovery import KnowledgeIngestRecoveryWorker
from control_plane.settings import Settings
from expert_work.persistence import InMemoryKnowledgeStore
from expert_work.protocol import (
    DEFAULT_CHUNK_MAX_TOKENS,
    DEFAULT_CHUNK_OVERLAP_TOKENS,
    DocumentPage,
    DocumentStatus,
    IngestErrorCode,
    IngestStage,
    PageStatus,
)
from expert_work.runtime.storage import InMemoryObjectStore
from orchestrator.llm import FakeEmbedder
from tests.figure_test_images import noise_png
from tests.knowledge_form_fixtures import picture_docx, picture_pptx

_TENANT = uuid4()

#: The refusal as stored — the console reads its two numbers
#: (``apps/admin-ui/src/pages/knowledge_detail/documentError.ts``, whose test
#: pins this same text).
_REFUSAL = (
    "text too large: 12,345,678 characters, over the per-document limit of 2,500,000; "
    "split the document into smaller files and upload them separately"
)


@pytest.fixture(autouse=True)
def _default_cap() -> Iterator[None]:
    yield
    configure_parsed_size(max_chars=DEFAULT_MAX_PARSED_CHARS)
    configure_raw_bytes(budget=DEFAULT_RAW_BYTES_BUDGET, upload_max_bytes=DEFAULT_UPLOAD_MAX_BYTES)


@pytest.fixture(scope="module", autouse=True)
def _parse_pool_teardown() -> Iterator[None]:
    yield
    shutdown_parse_pool()


# ---------------------------------------------------------------------------
# The cap
# ---------------------------------------------------------------------------


def test_the_default_is_the_settings_default() -> None:
    assert Settings().knowledge_document_max_chars == DEFAULT_MAX_PARSED_CHARS
    assert parsed_size.max_parsed_chars() == DEFAULT_MAX_PARSED_CHARS


@pytest.mark.parametrize(
    ("chars", "refused"),
    [(1000, False), (1001, True), (50_000_000, True)],
    ids=["at-the-limit", "just-over", "far-over"],
)
def test_the_cap_refuses_only_text_over_it(chars: int, refused: bool) -> None:
    if refused:
        with pytest.raises(DocumentTooLargeError) as caught:
            ensure_within_cap(chars, 1000)
        assert (caught.value.chars, caught.value.limit) == (chars, 1000)
    else:
        ensure_within_cap(chars, 1000)


@pytest.mark.parametrize(
    ("chars", "refused"),
    [(1000, False), (1001, True), (2_000_000, True)],
    ids=["at-the-limit", "just-over", "far-over"],
)
def test_parse_within_cap_measures_the_parsed_text(chars: int, refused: bool) -> None:
    raw = b"x" * chars
    if refused:
        with pytest.raises(DocumentTooLargeError) as caught:
            parse_within_cap("notes.txt", raw, 1000)
        assert caught.value.chars == chars
    else:
        assert parse_within_cap("notes.txt", raw, 1000) == "x" * chars


def test_the_refusal_says_how_large_what_the_limit_is_and_to_split() -> None:
    exc = DocumentTooLargeError(12_345_678, 2_500_000)
    assert str(exc) == too_large_message(12_345_678, 2_500_000) == _REFUSAL
    assert failure_code(exc) is IngestErrorCode.TEXT_TOO_LARGE
    match = TOO_LARGE_PATTERN.search(_REFUSAL[: ingestion._ERROR_CAP])
    assert match is not None and match.groups() == ("12,345,678", "2,500,000")
    # It is raised in the parse worker and travels back by pickle.
    back = pickle.loads(pickle.dumps(exc))  # noqa: S301 — our own exception
    assert (type(back), str(back), back.code, back.chars, back.limit) == (
        DocumentTooLargeError,
        _REFUSAL,
        IngestErrorCode.TEXT_TOO_LARGE,
        12_345_678,
        2_500_000,
    )


async def test_the_parse_worker_refuses_before_the_text_comes_back() -> None:
    configure_parsed_size(max_chars=1000)
    with pytest.raises(DocumentTooLargeError) as caught:
        await parse_document_in_pool("export.txt", b"y" * 5000)
    assert (caught.value.chars, caught.value.limit) == (5000, 1000)
    # Raised inside the worker (the pool attaches the worker's traceback), so
    # the 5,000 characters never crossed into this process.
    assert "parse_within_cap" in str(caught.value.__cause__)
    assert await parse_document_in_pool("export.txt", b"y" * 1000) == "y" * 1000


# ---------------------------------------------------------------------------
# Every format: the cap sits after parsing, in the worker
# ---------------------------------------------------------------------------

_PARAGRAPHS = [
    f"第{i}条:报销单据在出差结束后十个工作日内提交,由部门负责人审核后交财务复核。"
    for i in range(12)
]
_ENGLISH = [f"Item {i}: requests are approved by the team lead within two days." for i in range(12)]


def _txt() -> bytes:
    return "\n\n".join(_PARAGRAPHS).encode()


def _md() -> bytes:
    return ("# 报销制度\n\n## 提交\n\n" + "\n\n".join(_PARAGRAPHS)).encode()


def _html() -> bytes:
    body = "".join(f"<p>{p}</p>" for p in _PARAGRAPHS)
    return f"<html><body><h1>报销制度</h1>{body}</body></html>".encode()


def _csv() -> bytes:
    rows = ["编号,事项,说明", *(f"{i},报销,{p}" for i, p in enumerate(_PARAGRAPHS))]
    return "\n".join(rows).encode()


_DOCX_TYPES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" '
    'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" ContentType="application/'
    'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
)
_DOCX_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
    'relationships/officeDocument" Target="word/document.xml"/></Relationships>'
)


def _docx() -> bytes:
    body = "".join(f"<w:p><w:r><w:t>{escape(p)}</w:t></w:r></w:p>" for p in _PARAGRAPHS)
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document '
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body}</w:body></w:document>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr("[Content_Types].xml", _DOCX_TYPES)
        zf.writestr("_rels/.rels", _DOCX_RELS)
        zf.writestr("word/document.xml", document)
    return buffer.getvalue()


def _xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["编号", "说明"])
    for i, paragraph in enumerate(_PARAGRAPHS):
        sheet.append([i, paragraph])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pptx() -> bytes:
    deck = Presentation()
    for paragraph in _PARAGRAPHS:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
        box.text_frame.text = paragraph
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _pdf() -> bytes:
    doc = pymupdf.open()
    for start in range(0, len(_ENGLISH), 4):
        page = doc.new_page()
        page.insert_text((72, 72), "\n".join(_ENGLISH[start : start + 4]), fontsize=10)
    return bytes(doc.tobytes())


_FORMATS: list[tuple[str, Callable[[], bytes]]] = [
    ("notes.txt", _txt),
    ("policy.md", _md),
    ("policy.markdown", _md),
    ("policy.html", _html),
    ("policy.htm", _html),
    ("items.csv", _csv),
    ("policy.docx", _docx),
    ("items.xlsx", _xlsx),
    ("slides.pptx", _pptx),
    ("handbook.pdf", _pdf),
]


async def _document(store: InMemoryKnowledgeStore, filename: str, raw: bytes) -> tuple[UUID, UUID]:
    base = await store.create_base(tenant_id=_TENANT, name=f"kb-{uuid4().hex[:6]}")
    doc = await store.upsert_document(
        tenant_id=_TENANT, kb_id=base.id, filename=filename, content=raw
    )
    return base.id, doc.id


async def _ingest(
    store: InMemoryKnowledgeStore, kb_id: UUID, doc_id: UUID, filename: str, raw: bytes, **kw: Any
) -> int:
    return await ingest_document_bytes(
        store=store,
        embedder=FakeEmbedder(),
        tenant_id=_TENANT,
        document_id=doc_id,
        kb_id=kb_id,
        filename=filename,
        raw=raw,
        chunk_max_tokens=DEFAULT_CHUNK_MAX_TOKENS,
        chunk_overlap_tokens=DEFAULT_CHUNK_OVERLAP_TOKENS,
        **kw,
    )


@pytest.mark.parametrize(("filename", "build"), _FORMATS, ids=[name for name, _ in _FORMATS])
async def test_every_format_is_refused_in_the_worker_just_over_the_cap(
    filename: str, build: Callable[[], bytes], monkeypatch: pytest.MonkeyPatch
) -> None:
    raw = build()
    chars = len(parse_document(filename, raw))
    assert chars > 200  # a real document's worth of text, not an empty parse
    server_checks: list[int] = []
    real_check = ingestion.check_parsed_size

    def spy(text: str) -> None:
        server_checks.append(len(text))
        real_check(text)

    monkeypatch.setattr(ingestion, "check_parsed_size", spy)
    store = InMemoryKnowledgeStore()
    kb_id, doc_id = await _document(store, filename, raw)

    configure_parsed_size(max_chars=chars - 1)
    with pytest.raises(DocumentTooLargeError) as caught:
        await _ingest(store, kb_id, doc_id, filename, raw)
    assert (caught.value.chars, caught.value.limit) == (chars, chars - 1)
    assert server_checks == []  # refused in the worker: no text reached the server

    configure_parsed_size(max_chars=chars)
    assert await _ingest(store, kb_id, doc_id, filename, raw) > 0  # at the cap it ingests
    assert server_checks == [chars]


# Scanned pages: recognition stops paying once the document's text limit is
# reached, and what it read always fits — the document ingests.

_RECOGNIZED = "\n\n".join(_PARAGRAPHS)  # one page's reading
_TEXT_PAGE = "Leave requests are submitted three working days ahead and approved. " * 4


@dataclass
class _Settings:
    async def recognition_settings(self) -> RecognitionSettings:
        return RecognitionSettings(
            primary=ModelChoice("glm", "glm-ocr"), fallback=None, max_pages=1000
        )


class _Reader:
    """Fake recognition model; every call is a paid call."""

    def __init__(self) -> None:
        self.calls = 0

    async def read(self, image: bytes) -> Reading:
        del image
        self.calls += 1
        return Reading(text=_RECOGNIZED)


async def _key(*, tenant_id: UUID, provider: str) -> str:
    del tenant_id
    return f"key-{provider}"


def _scan(number: int) -> bytes:
    # A different picture on every page: identical ones would be read once and
    # then come from the document's recognition cache.
    source = pymupdf.open()
    picture = source.new_page()
    picture.insert_text((72, 72), f"Scanned page {number} " * 6, fontsize=14)
    return bytes(picture.get_pixmap(dpi=72).tobytes("png"))


def _pdf_pages(*kinds: str) -> bytes:
    doc = pymupdf.open()
    for number, kind in enumerate(kinds):
        page = doc.new_page()
        if kind == "text":
            page.insert_text((72, 72), _TEXT_PAGE, fontsize=8)
        else:
            page.insert_image(page.rect, stream=_scan(number))
    return bytes(doc.tobytes())


def _recognizer(store: InMemoryKnowledgeStore, reader: _Reader) -> DocumentRecognizer:
    def factory(choice: ModelChoice, api_key: str) -> RecognitionClient:
        del choice, api_key
        return reader

    return DocumentRecognizer(
        store=store, settings=_Settings(), api_keys=_key, usage=None, client_factory=factory
    )


@pytest.fixture
def _one_page_batches(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(recognizer_module, "RENDER_BATCH", 1)


async def _ingest_pdf(raw: bytes, reader: _Reader) -> tuple[int, list[DocumentPage]]:
    store = InMemoryKnowledgeStore()
    kb_id, doc_id = await _document(store, "scan.pdf", raw)
    count = await _ingest(
        store, kb_id, doc_id, "scan.pdf", raw, recognizer=_recognizer(store, reader)
    )
    pages = await store.list_document_pages(tenant_id=_TENANT, document_id=doc_id)
    assert pages is not None
    return count, pages


@pytest.mark.usefixtures("_one_page_batches")
async def test_recognition_stops_paying_once_the_text_limit_is_reached() -> None:
    # Room for exactly two pages' readings (plus the blank line joining each
    # of the five pages): pages three to five are never sent to the model.
    configure_parsed_size(max_chars=2 * len(_RECOGNIZED) + 2 * 5)
    reader = _Reader()
    count, pages = await _ingest_pdf(_pdf_pages(*["scan"] * 5), reader)
    assert reader.calls == 2
    assert [p.status for p in pages] == [PageStatus.RECOGNIZED] * 2 + [PageStatus.FAILED] * 3
    assert {p.error for p in pages[2:]} == {"document_text_limit"}
    assert count > 0  # the document still ingests, with the two pages read


@pytest.mark.usefixtures("_one_page_batches")
async def test_the_text_pages_count_against_the_recognition_limit() -> None:
    # A text page's own text plus one page's reading fill the cap: the second
    # scanned page is not read, and the assembled document is within the cap.
    text = len(_TEXT_PAGE.strip())
    configure_parsed_size(max_chars=text + len(_RECOGNIZED) + 2 * 3)
    reader = _Reader()
    count, pages = await _ingest_pdf(_pdf_pages("text", "scan", "scan"), reader)
    assert reader.calls == 1
    assert [p.status for p in pages[1:]] == [PageStatus.RECOGNIZED, PageStatus.FAILED]
    assert pages[2].error == "document_text_limit"
    assert count > 0


async def test_text_from_outside_a_parse_worker_is_checked_in_the_server() -> None:
    # Pre-parsed Markdown (and anything a later step adds) never went through
    # a worker's check: the server refuses it before chunking.
    configure_parsed_size(max_chars=1000)
    store = InMemoryKnowledgeStore()
    kb_id, doc_id = await _document(store, "notes.md", b"x")
    with pytest.raises(DocumentTooLargeError) as caught:
        await _ingest(store, kb_id, doc_id, "notes.md", b"x", parsed_markdown="y" * 1001)
    assert caught.value.chars == 1001
    doc = await store.get_document(tenant_id=_TENANT, document_id=doc_id)
    assert doc is not None and doc.stage is not IngestStage.CHUNKING


# Pictures described into the text (B-209 PR-2): the cap is on the text that
# reaches chunking, descriptions included. The worker refuses its own output
# (which the descriptions can only lengthen) before any picture is described;
# the server checks the text again once the descriptions are spliced in.

_FIGURE_DESCRIPTION = "一台设备的外观照片。"


class _Vision:
    """Fake picture-description model; every call is a paid call."""

    def __init__(self) -> None:
        self.calls = 0

    def factory(
        self, choice: ModelChoice, api_key: str, prompt: str, user_text: str, max_tokens: int
    ) -> _Vision:
        del choice, api_key, prompt, user_text, max_tokens
        return self

    async def read(self, image: bytes) -> Reading:
        del image
        self.calls += 1
        answer = {"text": "none", "chart": False, "description": _FIGURE_DESCRIPTION}
        return Reading(text=json.dumps(answer), finish_reason="stop", usage_reported=True)


@dataclass
class _FigureSettings:
    async def recognition_settings(self) -> RecognitionSettings:
        return RecognitionSettings(
            primary=ModelChoice("glm", "glm-ocr"),
            fallback=None,
            max_pages=1000,
            figure=ModelChoice("qwen", "qwen3.8-flash"),
        )


def _describer(store: InMemoryKnowledgeStore, vision: _Vision) -> FigureDescriber:
    settings = _FigureSettings()
    return FigureDescriber(
        store=store,
        settings=settings,
        api_keys=_key,
        recognizer=_recognizer(store, _Reader()),
        storage=FigureStorage(InMemoryObjectStore()),
        client_factory=vision.factory,
    )


def _manual_docx() -> bytes:
    return picture_docx(
        noise_png(640, 480), before="安装前请阅读本说明。" * 30, after="以上为整机外观。" * 30
    )


def _manual_pptx() -> bytes:
    return picture_pptx(noise_png(640, 480), title="整机外观与安装说明 " * 20)


@pytest.mark.parametrize(
    ("filename", "build"),
    [("manual.docx", _manual_docx), ("slides.pptx", _manual_pptx)],
    ids=["docx", "pptx"],
)
async def test_a_file_with_pictures_is_refused_before_any_is_described(
    filename: str, build: Callable[[], bytes], monkeypatch: pytest.MonkeyPatch
) -> None:
    raw = build()
    chars = len(parse_office_figures(raw, Path(filename).suffix).markdown)
    server_checks: list[int] = []
    real_check = ingestion.check_parsed_size

    def spy(text: str) -> None:
        server_checks.append(len(text))
        real_check(text)

    monkeypatch.setattr(ingestion, "check_parsed_size", spy)
    store = InMemoryKnowledgeStore()
    kb_id, doc_id = await _document(store, filename, raw)
    vision = _Vision()
    configure_parsed_size(max_chars=chars - 1)
    with pytest.raises(DocumentTooLargeError) as caught:
        await _ingest(store, kb_id, doc_id, filename, raw, figures=_describer(store, vision))
    assert caught.value.chars == chars
    assert (vision.calls, server_checks) == (0, [])  # refused in the worker, nothing paid


async def test_picture_descriptions_count_against_the_cap() -> None:
    # The file's own text fits; with its picture's description in place it
    # does not — the server refuses it before chunking. With room for the
    # description it ingests.
    raw = _manual_docx()
    chars = len(parse_office_figures(raw, ".docx").markdown)
    store = InMemoryKnowledgeStore()
    kb_id, doc_id = await _document(store, "manual.docx", raw)
    vision = _Vision()
    figures = _describer(store, vision)
    configure_parsed_size(max_chars=chars)
    with pytest.raises(DocumentTooLargeError) as caught:
        await _ingest(store, kb_id, doc_id, "manual.docx", raw, figures=figures)
    assert vision.calls == 1 and caught.value.chars > chars
    configure_parsed_size(max_chars=caught.value.chars)
    assert await _ingest(store, kb_id, doc_id, "manual.docx", raw, figures=figures) > 0


# ---------------------------------------------------------------------------
# A refused document: failed for good, the rest of the batch unaffected
# ---------------------------------------------------------------------------


async def _two_documents(store: InMemoryKnowledgeStore) -> tuple[UUID, UUID]:
    base = await store.create_base(tenant_id=_TENANT, name=f"kb-{uuid4().hex[:6]}")
    big = await store.upsert_document(
        tenant_id=_TENANT, kb_id=base.id, filename="export.txt", content=b"z" * 5000
    )
    small = await store.upsert_document(
        tenant_id=_TENANT, kb_id=base.id, filename="policy.md", content=_md()
    )
    return big.id, small.id


async def test_a_refused_document_fails_once_and_the_rest_of_the_batch_ingests() -> None:
    configure_parsed_size(max_chars=1000)
    store = InMemoryKnowledgeStore()
    big_id, small_id = await _two_documents(store)
    runner = KnowledgeIngestionRunner(store=store, embedder=FakeEmbedder(), concurrency=4)
    for doc_id in (big_id, small_id):
        runner.submit(tenant_id=_TENANT, document_id=doc_id)
    await runner.drain()

    big = await store.get_document(tenant_id=_TENANT, document_id=big_id)
    small = await store.get_document(tenant_id=_TENANT, document_id=small_id)
    assert big is not None and small is not None
    assert big.status is DocumentStatus.FAILED
    assert big.error_code is IngestErrorCode.TEXT_TOO_LARGE
    assert big.error == too_large_message(5000, 1000)
    assert big.attempts == 1
    assert small.status is DocumentStatus.READY and small.chunk_count > 0

    # Terminal: the recovery sweep does not pick it up again.
    worker = _recovery(store)
    assert await worker.run_once() == 0
    again = await store.get_document(tenant_id=_TENANT, document_id=big_id)
    assert again is not None and again.attempts == 1 and again.status is DocumentStatus.FAILED


def _recovery(
    store: InMemoryKnowledgeStore, *, embedder: Any = None
) -> KnowledgeIngestRecoveryWorker:
    return KnowledgeIngestRecoveryWorker(
        store=store,
        embedder=embedder or FakeEmbedder(),
        interval_s=1,
        concurrency=4,
        lease_seconds=300,
        max_attempts=5,
    )


async def test_the_recovery_worker_fails_a_refused_document_on_the_first_attempt() -> None:
    # Retries remain (1 of 5) — but the text will be just as long next time:
    # failing it now spares four more full parses of an oversized document.
    configure_parsed_size(max_chars=1000)
    store = InMemoryKnowledgeStore()
    big_id, small_id = await _two_documents(store)
    assert await _recovery(store).run_once() == 2

    big = await store.get_document(tenant_id=_TENANT, document_id=big_id)
    small = await store.get_document(tenant_id=_TENANT, document_id=small_id)
    assert big is not None and small is not None
    assert (big.status, big.error_code, big.attempts) == (
        DocumentStatus.FAILED,
        IngestErrorCode.TEXT_TOO_LARGE,
        1,
    )
    assert small.status is DocumentStatus.READY


# ---------------------------------------------------------------------------
# The budget: what is embedded / stored at once
# ---------------------------------------------------------------------------


async def _settle() -> None:
    for _ in range(20):
        await asyncio.sleep(0)


async def test_the_budget_holds_a_document_until_the_text_in_flight_fits() -> None:
    configure_parsed_size(max_chars=1000)
    order: list[str] = []
    release = asyncio.Event()

    async def ingest(name: str, chars: int) -> None:
        async with admitted(chars):
            order.append(name)
            await release.wait()

    first = asyncio.create_task(ingest("first", 700))
    await _settle()
    second = asyncio.create_task(ingest("second", 700))
    await _settle()
    assert order == ["first"]  # 700 + 700 > 1000: the second waits
    release.set()
    await asyncio.gather(first, second)
    assert order == ["first", "second"]


async def test_the_budget_is_first_come_first_served() -> None:
    # A small document that would fit does not overtake a larger one waiting
    # ahead of it — otherwise a stream of small uploads starves a large one.
    configure_parsed_size(max_chars=1000)
    order: list[str] = []
    gates = {name: asyncio.Event() for name in ("a", "b", "c")}

    async def ingest(name: str, chars: int) -> None:
        async with admitted(chars):
            order.append(name)
            await gates[name].wait()

    tasks = []
    for name, chars in (("a", 600), ("b", 600), ("c", 300)):
        tasks.append(asyncio.create_task(ingest(name, chars)))
        await _settle()
    assert order == ["a"]
    gates["a"].set()
    await _settle()
    assert order == ["a", "b", "c"]  # b first, then c fits beside it
    for gate in gates.values():
        gate.set()
    await asyncio.gather(*tasks)


async def test_a_cancelled_waiter_lets_the_next_one_in() -> None:
    configure_parsed_size(max_chars=1000)
    order: list[str] = []
    hold = asyncio.Event()

    async def ingest(name: str, chars: int) -> None:
        async with admitted(chars):
            order.append(name)
            await hold.wait()

    first = asyncio.create_task(ingest("first", 600))
    await _settle()
    blocked = asyncio.create_task(ingest("blocked", 600))
    await _settle()
    behind = asyncio.create_task(ingest("behind", 300))
    await _settle()
    assert order == ["first"]
    blocked.cancel()  # the upload was deleted while it waited
    await _settle()
    assert order == ["first", "behind"]  # 600 + 300 fit once the head is gone
    hold.set()
    await asyncio.gather(first, behind)
    assert blocked.cancelled()


async def _fail_while_admitted(chars: int) -> None:
    async with admitted(chars):
        raise RuntimeError("embedding failed")


async def test_the_budget_is_returned_when_an_ingest_fails() -> None:
    configure_parsed_size(max_chars=1000)
    with pytest.raises(RuntimeError):
        await _fail_while_admitted(1000)
    async with asyncio.timeout(1):
        async with admitted(1000):
            pass


class _GatedEmbedder:
    """FakeEmbedder whose first call for chunks (their texts carry the
    document title; topic detection embeds bare blocks) waits for ``gate``."""

    def __init__(self) -> None:
        self.gate = asyncio.Event()
        self._inner = FakeEmbedder()
        self._first = True

    async def embed(self, texts: Sequence[str], *, tenant_id: UUID) -> list[tuple[float, ...]]:
        if self._first and texts and texts[0].startswith("[Document: "):
            self._first = False
            await self.gate.wait()
        return await self._inner.embed(texts, tenant_id=tenant_id)


async def test_concurrent_ingests_wait_for_the_budget_after_chunking(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    store = InMemoryKnowledgeStore()
    base = await store.create_base(tenant_id=_TENANT, name="kb-budget")
    text = "\n\n".join(_PARAGRAPHS).encode()
    ids = [
        (
            await store.upsert_document(
                tenant_id=_TENANT, kb_id=base.id, filename=f"part-{i}.txt", content=text
            )
        ).id
        for i in range(2)
    ]
    configure_parsed_size(max_chars=len(text.decode()) + 10)  # one fits, two do not
    waiting: list[int] = []
    inside: list[int] = []
    real = ingestion.admitted

    @asynccontextmanager
    async def spy(chars: int, chunks: int = 0) -> AsyncIterator[None]:
        waiting.append(chars)
        async with real(chars, chunks):
            inside.append(chars)
            yield

    monkeypatch.setattr(ingestion, "admitted", spy)
    embedder = _GatedEmbedder()
    runner = KnowledgeIngestionRunner(store=store, embedder=embedder, concurrency=4)
    for doc_id in ids:
        runner.submit(tenant_id=_TENANT, document_id=doc_id)
    async with asyncio.timeout(60):
        while len(waiting) < 2:
            await asyncio.sleep(0.01)
    await _settle()
    # Both chunked (outside the budget: the text and a window of block
    # vectors); only one is embedded / stored at a time (B-214).
    assert len(inside) == 1
    documents = [await store.get_document(tenant_id=_TENANT, document_id=i) for i in ids]
    assert {doc.stage for doc in documents if doc is not None} == {
        IngestStage.EMBEDDING,
        IngestStage.CHUNKING,
    }
    embedder.gate.set()
    await runner.drain()
    assert len(inside) == 2
    for doc_id in ids:
        doc = await store.get_document(tenant_id=_TENANT, document_id=doc_id)
        assert doc is not None and doc.status is DocumentStatus.READY


# ---------------------------------------------------------------------------
# The raw-bytes budget: original upload bytes held at once
# ---------------------------------------------------------------------------


def test_the_raw_bytes_defaults_are_the_settings_defaults() -> None:
    settings = Settings()
    assert settings.knowledge_ingest_raw_bytes_budget == DEFAULT_RAW_BYTES_BUDGET
    assert settings.knowledge_document_max_bytes == DEFAULT_UPLOAD_MAX_BYTES


async def test_a_claim_reserves_the_largest_upload_and_keeps_its_real_size() -> None:
    configure_raw_bytes(budget=100, upload_max_bytes=60)
    first = await reserve_raw_bytes()  # 60 held: a second 60 does not fit
    second = asyncio.create_task(reserve_raw_bytes())
    await _settle()
    assert not second.done()
    first.shrink(30)  # the bytes read were 30: 30 + 60 fit
    held = await asyncio.wait_for(second, 1)
    held.release()
    first.release()


class _ParkingEmbedder(FakeEmbedder):
    """Every call waits for ``gate`` — documents stay in flight until then."""

    def __init__(self) -> None:
        super().__init__()
        self.gate = asyncio.Event()

    async def embed(self, texts: Sequence[str], *, tenant_id: UUID) -> list[tuple[float, ...]]:
        await self.gate.wait()
        return await super().embed(texts, tenant_id=tenant_id)


async def _uploads(store: InMemoryKnowledgeStore, size: int, count: int) -> list[UUID]:
    base = await store.create_base(tenant_id=_TENANT, name=f"kb-{uuid4().hex[:6]}")
    ids = []
    for i in range(count):
        body = f"# Part {i}\n\n" + "x" * size
        doc = await store.upsert_document(
            tenant_id=_TENANT, kb_id=base.id, filename=f"part-{i}.md", content=body.encode()[:size]
        )
        ids.append(doc.id)
    return ids


async def _statuses(store: InMemoryKnowledgeStore, ids: Sequence[UUID]) -> list[DocumentStatus]:
    docs = [await store.get_document(tenant_id=_TENANT, document_id=i) for i in ids]
    return [doc.status for doc in docs if doc is not None]


async def _until(predicate: Callable[[], Any]) -> None:
    async with asyncio.timeout(15):
        while not await predicate():
            await asyncio.sleep(0.01)


async def test_uploads_wait_for_room_before_their_bytes_are_read() -> None:
    # Room for one 55-byte upload once its 60-byte reservation shrinks; the
    # second waits *before* its claim reads the bytes, so it stays pending.
    configure_raw_bytes(budget=100, upload_max_bytes=60)
    store = InMemoryKnowledgeStore()
    ids = await _uploads(store, 55, 2)
    embedder = _ParkingEmbedder()
    runner = KnowledgeIngestionRunner(store=store, embedder=embedder, concurrency=4)
    for doc_id in ids:
        runner.submit(tenant_id=_TENANT, document_id=doc_id)

    async def one_in_flight() -> bool:
        return (await _statuses(store, ids)).count(DocumentStatus.PROCESSING) == 1

    await _until(one_in_flight)
    await asyncio.sleep(0.2)
    assert sorted(await _statuses(store, ids)) == sorted(
        [DocumentStatus.PROCESSING, DocumentStatus.PENDING]
    )
    embedder.gate.set()
    await asyncio.wait_for(runner.drain(), 15)
    assert await _statuses(store, ids) == [DocumentStatus.READY, DocumentStatus.READY]


async def test_small_uploads_share_the_room_a_large_reservation_leaves() -> None:
    # 30-byte uploads keep 30 of their 60-byte reservation: both are in flight.
    configure_raw_bytes(budget=100, upload_max_bytes=60)
    store = InMemoryKnowledgeStore()
    ids = await _uploads(store, 30, 2)
    embedder = _ParkingEmbedder()
    runner = KnowledgeIngestionRunner(store=store, embedder=embedder, concurrency=4)
    for doc_id in ids:
        runner.submit(tenant_id=_TENANT, document_id=doc_id)

    async def both_in_flight() -> bool:
        return (await _statuses(store, ids)).count(DocumentStatus.PROCESSING) == 2

    await _until(both_in_flight)
    embedder.gate.set()
    await asyncio.wait_for(runner.drain(), 15)


async def test_the_recovery_worker_claims_within_the_same_budget() -> None:
    # An upload in flight holds 55 bytes; the recovery sweep's claim waits for
    # room instead of reading another document's bytes, then drives it.
    configure_raw_bytes(budget=100, upload_max_bytes=60)
    store = InMemoryKnowledgeStore()
    uploaded, stuck = await _uploads(store, 55, 2)
    embedder = _ParkingEmbedder()
    runner = KnowledgeIngestionRunner(store=store, embedder=embedder, concurrency=4)
    runner.submit(tenant_id=_TENANT, document_id=uploaded)

    async def uploaded_in_flight() -> bool:
        return (await _statuses(store, [uploaded])) == [DocumentStatus.PROCESSING]

    await _until(uploaded_in_flight)
    cycle = asyncio.create_task(_recovery(store, embedder=embedder).run_once())
    await asyncio.sleep(0.2)
    assert await _statuses(store, [stuck]) == [DocumentStatus.PENDING]
    embedder.gate.set()
    assert await asyncio.wait_for(cycle, 15) == 1
    await asyncio.wait_for(runner.drain(), 15)
    assert await _statuses(store, [uploaded, stuck]) == [DocumentStatus.READY] * 2
