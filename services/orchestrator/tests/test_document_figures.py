"""``document_figures`` —— 来源侧图清单。

两层,与 test_read_document 同构:
  1. 沙箱内探测片段,拿本地 temp 工作区直接跑;
  2. 信封 → 清单文本的渲染。
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest

from orchestrator.tools.document_figures import (
    MIN_FIGURE_EDGE_PT,
    build_figure_inventory_wrapper,
    render_figure_map,
)


def _run(tmp_path: Path, rel: str) -> dict:
    """本地执行沙箱片段(与 test_read_document 的做法一致)。"""
    code = build_figure_inventory_wrapper(rel, ws=str(tmp_path), max_bytes=25 * 1024 * 1024)
    ns: dict = {}
    import contextlib

    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        exec(compile(code, "<snippet>", "exec"), ns)  # noqa: S102
    return json.loads(buf.getvalue())


def _png(w: int, h: int) -> io.BytesIO:
    Image = pytest.importorskip("PIL.Image")  # noqa: N806
    b = io.BytesIO()
    Image.new("RGB", (w, h), "#cde").save(b, "PNG")
    b.seek(0)
    return b


def test_pptx_picture_slide_is_listed(tmp_path: Path) -> None:
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    prs = pptx.Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[5])
    s1.shapes.title.text = "封面"
    s1.shapes.add_picture(_png(800, 500), Inches(1), Inches(2), width=Inches(5))
    prs.save(tmp_path / "d.pptx")

    env = _run(tmp_path, "d.pptx")
    assert env["ok"] is True
    assert env["state"] == "figures"
    assert env["format"] == "pptx"
    assert [f["unit"] for f in env["figures"]] == [1]
    assert env["figures"][0]["kind"] == "picture"
    assert env["figures"][0]["title"] == "封面"


def test_pptx_native_chart_becomes_data_not_a_figure(tmp_path: Path) -> None:
    """图表是数据,不该逼模型去看图。"""
    pptx = pytest.importorskip("pptx")
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "趋势"
    cd = CategoryChartData()
    cd.categories = ["7月", "8月", "9月"]
    cd.add_series("血糖", (6.1, 5.8, 5.5))
    s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(1), Inches(1.5), Inches(6), Inches(4), cd)
    prs.save(tmp_path / "c.pptx")

    env = _run(tmp_path, "c.pptx")
    chart = next(f for f in env["figures"] if f["kind"] == "chart")
    assert chart["chart_data"]["cats"] == ["7月", "8月", "9月"]
    assert chart["chart_data"]["series"] == {"血糖": [6.1, 5.8, 5.5]}


def test_decorative_logo_is_skipped(tmp_path: Path) -> None:
    """页脚 logo 29x29pt 占页 0.17%,不进清单;同页的 288x192pt 图表进。"""
    docx = pytest.importorskip("docx")
    from docx.shared import Inches

    d = docx.Document()
    d.add_heading("一、体检概览", 1)
    d.add_paragraph("这是第一段正文，后面跟着一张血糖趋势图。")  # noqa: RUF001
    d.add_picture(_png(600, 400), width=Inches(4))
    d.add_paragraph("下面是一张小 logo。")
    d.add_picture(_png(40, 40), width=Inches(0.4))
    d.save(tmp_path / "d.docx")

    env = _run(tmp_path, "d.docx")
    assert len(env["figures"]) == 1
    assert env["skipped_decorative"] == 1
    assert env["figures"][0]["w_pt"] == pytest.approx(288.0, abs=1.0)
    assert min(env["figures"][0]["w_pt"], env["figures"][0]["h_pt"]) >= MIN_FIGURE_EDGE_PT


def test_docx_anchor_is_the_sentence_before_the_figure(tmp_path: Path) -> None:
    docx = pytest.importorskip("docx")
    from docx.shared import Inches

    d = docx.Document()
    d.add_heading("一、体检概览", 1)
    d.add_paragraph("这是第一段正文，后面跟着一张血糖趋势图。")  # noqa: RUF001
    d.add_picture(_png(600, 400), width=Inches(4))
    d.save(tmp_path / "a.docx")

    env = _run(tmp_path, "a.docx")
    assert "血糖趋势图" in env["figures"][0]["anchor"]


def test_plain_table_is_not_a_figure(tmp_path: Path) -> None:
    """纯表格零误报 —— 表格不是 drawing,没有 graphicData。"""
    docx = pytest.importorskip("docx")

    d = docx.Document()
    d.add_heading("表格页", 1)
    t = d.add_table(rows=6, cols=4)
    for r in range(6):
        for c in range(4):
            t.cell(r, c).text = f"{r}-{c}"
    d.save(tmp_path / "t.docx")

    env = _run(tmp_path, "t.docx")
    assert env["state"] == "none"
    assert env["figures"] == []


def test_xlsx_never_asks_for_pixels(tmp_path: Path) -> None:
    """xlsx 的图表是单元格引用;不渲染,只报引用。"""
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.chart import LineChart, Reference

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "体检"
    for row in (["月", "血糖"], ["7月", 6.1], ["8月", 5.8], ["9月", 5.5]):
        ws.append(row)
    ch = LineChart()
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    ws.add_chart(ch, "D2")
    wb.save(tmp_path / "s.xlsx")

    env = _run(tmp_path, "s.xlsx")
    assert all(f["kind"] != "picture" for f in env["figures"])
    chart = next(f for f in env["figures"] if f["kind"] == "chart")
    assert "体检" in json.dumps(chart["chart_data"], ensure_ascii=False)


def test_corrupt_file_is_undetermined_not_none(tmp_path: Path) -> None:
    """「测不了」必须与「没有图」分开 —— 三态的核心。"""
    (tmp_path / "bad.pptx").write_bytes(b"not a zip at all")
    env = _run(tmp_path, "bad.pptx")
    assert env["state"] == "undetermined"
    assert env["state"] != "none"


# --- 清单文本渲染 -----------------------------------------------------------


def test_map_says_it_could_not_tell() -> None:
    text = render_figure_map({"ok": True, "state": "undetermined", "reason": "parse_failed"})
    assert "无法确定" in text


def test_map_is_empty_when_no_figures() -> None:
    assert render_figure_map({"ok": True, "state": "none", "figures": []}) == ""


def test_map_forbids_rendering_everything() -> None:
    env = {
        "ok": True,
        "state": "figures",
        "format": "pptx",
        "figures": [
            {
                "unit": 3,
                "kind": "picture",
                "count": 1,
                "w_pt": 288.0,
                "h_pt": 192.0,
                "anchor": "趋势",
                "alt": None,
                "title": "趋势",
            }
        ],
        "skipped_decorative": 0,
    }
    text = render_figure_map(env)
    assert "不要把所有" in text
    assert "read_page" in text
    assert "OCR" not in text.upper() or "技能" in text  # 不点名具体 skill


def test_map_caps_entries() -> None:
    env = {
        "ok": True,
        "state": "figures",
        "format": "pptx",
        "skipped_decorative": 0,
        "figures": [
            {
                "unit": i,
                "kind": "picture",
                "count": 1,
                "w_pt": 288.0,
                "h_pt": 192.0,
                "anchor": "",
                "alt": None,
                "title": None,
            }
            for i in range(1, 31)
        ],
    }
    text = render_figure_map(env)
    assert text.count("\n  ") <= 22  # 20 条 + 折叠行的余量
    assert "另有" in text
