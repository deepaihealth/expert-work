"""`kb_build_dataset.py`(B-175 L1-T2)的单元测试:全部用 tmp_path 造的小文件,不碰真语料。"""

# ruff: noqa: RUF001  # 中文标点是要匹配 / 要测的目标
from __future__ import annotations

import csv
from pathlib import Path

import openpyxl  # type: ignore[import-untyped]
import pytest
import yaml
from kb_build_dataset import (  # type: ignore[import-not-found]
    QaRecord,
    _parse_cached,
    build_dataset,
    contains_personal_data,
    dedupe_records,
    extract_docx_qa,
    extract_xlsx_qa,
    normalize_query,
    pick_expected,
)
from kb_dataset import load_dataset  # type: ignore[import-not-found]

_DOCX_MD = """# 产品说明书

## 七、产品常见QandA

1. **晚安饮什么时候吃？**

**答：**睡前30分钟服用最佳。研究指出口服后30-60分钟浓度达峰，利于睡眠起始。

2. **会不会有依赖性？**

**答：**无化学添加，成分天然。

显示在适量下无依赖风险。

## 八、其它

不是问答。
"""


def _write_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "通用"
    ws.append(["问题", "参考回复"])
    ws.append(["益生菌吃了会有依赖性吗?", "不会。正常情况下可通过饮食获得益生菌。"])
    ws.append(["益生菌吃了会有依赖性吗?", "华大益生菌均属于功能性食品，长期食用不会产生依赖性。"])
    ws2 = wb.create_sheet("甲产品")
    ws2.append([None, "问题", "回答", "备注"])
    ws2.append(["甲产品", "产品作用？", "适合日常餐后食用，改善肠道环境。", None])
    ws2.append([None, "怎么保存？", "常温密封保存即可，开袋后尽快食用。", "备注忽略"])
    ws3 = wb.create_sheet("跳过我")
    ws3.append(["问题搜集"])
    wb.save(path)


def test_extract_xlsx_qa_finds_header_row_per_sheet_and_tags_by_sheet(tmp_path: Path) -> None:
    path = tmp_path / "faq.xlsx"
    _write_xlsx(path)
    records = extract_xlsx_qa(
        path,
        source="faq_2026",
        question_headers=("问题",),
        answer_headers=("参考回复", "回答"),
        skip_sheets=("跳过我",),
    )
    # 产品分表里的短问题补了产品名作主语;通用表的不补
    assert [r.query for r in records] == [
        "益生菌吃了会有依赖性吗?",
        "益生菌吃了会有依赖性吗?",
        "甲产品产品作用？",
        "甲产品怎么保存？",
    ]
    assert records[0].tags == ("通用",)
    assert records[2].tags == ("甲产品",)
    assert records[3].answer == "常温密封保存即可，开袋后尽快食用。"
    assert all(r.source == "faq_2026" for r in records)
    assert all(r.origin_document == "faq.xlsx" for r in records)


def test_extract_xlsx_qa_can_take_tag_from_a_column(tmp_path: Path) -> None:
    path = tmp_path / "kb.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "产品FAQ"
    ws.append(["类型", "问题", "专业参考回复", "备注"])
    ws.append(["益生菌通用", "要冷藏吗？", "常温密封保存即可。", None])
    wb.save(path)
    records = extract_xlsx_qa(
        path,
        source="faq_xlsx",
        question_headers=("问题",),
        answer_headers=("专业参考回复",),
        tag_header="类型",
        sheets=("产品FAQ",),
    )
    assert len(records) == 1
    assert records[0].tags == ("益生菌通用",)


def test_extract_docx_qa_reads_numbered_bold_questions_until_next_heading() -> None:
    records = extract_docx_qa(
        _DOCX_MD, source="manual_qa", origin_document="说明书.docx", section_marker="常见Q"
    )
    assert [r.query for r in records] == ["晚安饮什么时候吃？", "会不会有依赖性？"]
    assert records[0].answer.startswith("睡前30分钟服用最佳。")
    assert "答：" not in records[0].answer
    # 多段答案合并,且不吞掉下一节
    assert "显示在适量下无依赖风险。" in records[1].answer
    assert "不是问答" not in records[1].answer


def test_normalize_query_ignores_numbering_punctuation_and_width() -> None:
    assert normalize_query("2. 益生菌需要放冰箱冷藏吗？") == normalize_query(
        "益生菌需要放冰箱冷藏吗?"
    )
    assert normalize_query("产品作用？") != normalize_query("产品成分？")


def _rec(
    query: str, answer: str, *, source: str = "faq_xlsx", tags: tuple[str, ...] = ()
) -> QaRecord:
    return QaRecord(query=query, answer=answer, source=source, tags=tags, origin_document="a.xlsx")


def test_dedupe_keeps_first_query_merges_tags_and_answers() -> None:
    merged = dedupe_records(
        [
            _rec("益生菌要冷藏吗？", "常温即可。", tags=("通用",)),
            _rec("1. 益生菌要冷藏吗?", "放阴凉处。", source="manual_qa", tags=("甲产品",)),
            _rec("每天怎么吃？", "温水冲服。"),
        ]
    )
    assert [m.query for m in merged] == ["益生菌要冷藏吗？", "每天怎么吃？"]
    assert merged[0].source == "faq_xlsx"
    assert merged[0].tags == ("通用", "甲产品")
    assert merged[0].answers == ("常温即可。", "放阴凉处。")


def test_dedupe_drops_queries_shorter_than_the_minimum() -> None:
    merged = dedupe_records([_rec("健康指导", "……"), _rec("益生菌要冷藏吗？", "常温即可。")])
    assert [m.query for m in merged] == ["益生菌要冷藏吗？"]


def test_contains_personal_data_flags_phone_wechat_email() -> None:
    assert contains_personal_data("请联系 13812345678 咨询")
    assert contains_personal_data("加微信 vx: healthy_2026 领取")
    assert contains_personal_data("邮件 someone@example.com")
    assert not contains_personal_data("专利号 ZL202011248842.0,每袋 2 克")


def test_pick_expected_prefers_origin_doc_then_other_docs_and_caps_at_three() -> None:
    parsed = {
        "说明书.docx": "……睡前30分钟服用最佳。……",
        "课件.pdf": "……睡前30分钟服用最佳。……",
        "faq.xlsx": "……睡前30分钟服用最佳。研究指出口服后30-60分钟浓度达峰……",
        "折页.pdf": "……睡前30分钟服用最佳。……",
        "无关.docx": "完全不相关的内容。",
    }
    expected = pick_expected(
        answers=("睡前30分钟服用最佳。研究指出口服后30-60分钟浓度达峰，利于睡眠起始。",),
        parsed_texts=parsed,
        origin_document="faq.xlsx",
        max_docs=3,
    )
    docs = [e.document for e in expected]
    assert docs[-1] == "faq.xlsx"  # 出处文档(答案原文所在)排最后,前面留给真正要找的文档
    assert len(docs) == 3
    assert "无关.docx" not in docs
    assert all(e.anchor in parsed[e.document] for e in expected)
    assert all(len(e.anchor) >= 8 for e in expected)


def test_pick_expected_origin_first_for_manual_qa() -> None:
    parsed = {
        "说明书.docx": "……餐前搭配膳食纤维使用效果更佳。……",
        "课件.pdf": "……餐前搭配膳食纤维使用效果更佳。……",
    }
    expected = pick_expected(
        answers=("餐前搭配膳食纤维使用效果更佳。",),
        parsed_texts=parsed,
        origin_document="说明书.docx",
        max_docs=3,
        origin_first=True,
    )
    assert [e.document for e in expected] == ["说明书.docx", "课件.pdf"]


def test_pick_expected_returns_empty_when_no_sentence_matches() -> None:
    assert (
        pick_expected(
            answers=("这句话哪里都没有出现过的样子。",),
            parsed_texts={"a.md": "别的内容"},
            origin_document=None,
            max_docs=3,
        )
        == ()
    )


def test_build_dataset_end_to_end_writes_loadable_yaml_and_review(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    (corpus / "甲产品").mkdir(parents=True)
    _write_xlsx(corpus / "faq.xlsx")
    (corpus / "甲产品" / "说明书.md").write_text(
        "# 甲产品\n\n常温密封保存即可，开袋后尽快食用。\n\n适合日常餐后食用，改善肠道环境。\n",
        encoding="utf-8",
    )
    (corpus / "无关.md").write_text("无关内容。\n", encoding="utf-8")
    manifest = tmp_path / "manifest.csv"
    with manifest.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(
            [
                "path",
                "top_dir",
                "ext",
                "size",
                "tier",
                "pages",
                "scanned_pages",
                "text_chars",
                "dup_group",
                "version_marker",
                "error",
            ]
        )
        # faq.xlsx 只做问题来源、不在语料清单里 → 只在它里面出现的答案配不到锚
        writer.writerow(["甲产品/说明书.md", "甲产品", "md", 1, "A", 0, 0, 0, "", "", ""])
        writer.writerow(["无关.md", ".", "md", 1, "A", 0, 0, 0, "", "", ""])
        writer.writerow(["扫描.pdf", ".", "pdf", 1, "B", 3, 3, 0, "", "", ""])
    extra = tmp_path / "hard.yaml"
    extra.write_text(
        yaml.safe_dump(
            {
                "cases": [
                    {
                        "id": "no_answer-0001",
                        "query": "有没有宠物版？",
                        "source": "no_answer",
                        "tags": [],
                        "expected": [],
                        "answer": None,
                    },
                    {
                        "id": "hard_exact-0001",
                        "query": "保存方式是什么？",
                        "source": "hard_exact",
                        "tags": [],
                        "expected": [
                            {
                                "document": "说明书.md",
                                "anchor": "常温密封保存即可，开袋后尽快食用",
                                "page": None,
                            },
                            {
                                "document": "说明书.md",
                                "anchor": "这句话在说明书里不存在",
                                "page": None,
                            },
                            {
                                "document": "不存在.md",
                                "anchor": "常温密封保存即可，开袋后尽快食用",
                                "page": None,
                            },
                        ],
                        "answer": None,
                    },
                ]
            },
            allow_unicode=True,
        ),
        encoding="utf-8",
    )
    spec = {
        "name": "demo_v1",
        "corpus": {"manifest": str(manifest), "tiers": ["A"]},
        "sources": [
            {
                "type": "xlsx_qa",
                "source": "faq_2026",
                "file": "faq.xlsx",
                "question_headers": ["问题"],
                "answer_headers": ["参考回复", "回答"],
                "skip_sheets": ["跳过我"],
            }
        ],
        "extra_cases": [str(extra)],
    }
    out = tmp_path / "demo_v1.yaml"
    review = tmp_path / "demo_v1.review.md"
    stats = build_dataset(
        spec,
        corpus_root=corpus,
        manifest=manifest,
        parsed_cache=tmp_path / "cache",
        out=out,
        review=review,
    )
    dataset = load_dataset(out)
    by_query = {c.query: c for c in dataset.cases}
    assert dataset.name == "demo_v1"
    # 两条配到了说明书;依赖性那条在语料里没有 → 进 review、不进评测集
    assert by_query["甲产品怎么保存？"].expected[0].document == "说明书.md"
    assert by_query["甲产品产品作用？"].expected[0].anchor in "适合日常餐后食用，改善肠道环境。"
    assert "益生菌吃了会有依赖性吗?" not in by_query
    assert "益生菌吃了会有依赖性吗?" in review.read_text(encoding="utf-8")
    assert by_query["有没有宠物版？"].expected == ()
    assert stats["matched"] == 2
    assert stats["unmatched"] == 1
    assert stats["extra"] == 2
    # 手写样本:锚文本和文档都逐条核对,不过的列进 review
    assert stats["extra_anchor_problems"] == 2
    review_text = review.read_text(encoding="utf-8")
    assert "hard_exact-0001: anchor not found in 说明书.md" in review_text
    assert "hard_exact-0001: document not in parsed corpus: 不存在.md" in review_text
    # 缓存落在仓库外的指定目录
    assert any((tmp_path / "cache").iterdir())


def test_parse_cached_strips_lone_surrogates_and_never_raises(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import control_plane.knowledge.parsing as parsing

    doc = tmp_path / "a.docx"
    doc.write_bytes(b"x")
    monkeypatch.setattr(parsing, "parse_document", lambda name, raw: "正文\ud83d尾巴")
    text = _parse_cached(doc, tmp_path / "cache")
    assert text is not None
    assert "\ud83d" not in text
    assert text.startswith("正文")
    assert (tmp_path / "cache").is_dir()
    monkeypatch.setattr(
        parsing, "parse_document", lambda name, raw: (_ for _ in ()).throw(RuntimeError("boom"))
    )
    assert _parse_cached(tmp_path / "b.docx", tmp_path / "cache") is None
